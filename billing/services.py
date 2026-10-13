"""
Billing klien (requirement §10): tagihan dari tb_invoice per nomor kontrak.

- Outstanding : tb_invoice.tgl_receive_payment kosong.
- History     : tb_invoice.tgl_receive_payment terisi (tanggal bayar dari klien).
  tb_invoice.tgl_pay tetap tanggal bayar internal SLI (dipakai saat generate invoice).
- Upload      : hasil export + kolom TANGGAL BAYAR; kunci NOMOR INVOICE + NOMOR KONTRAK.
  Partial payment per kontrak: baris yang tanggal bayarnya kosong tetap outstanding.
"""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from AktaSertifikat.models import LegalDocuments
from import_cabang.models import Consumer
from invoice.models import tb_invoice

COL_INVOICE = "NOMOR INVOICE"
COL_KONTRAK = "NOMOR KONTRAK"
COL_TGL_BAYAR = "TANGGAL BAYAR"

EXPORT_HEADERS = [
    COL_INVOICE, "TANGGAL INVOICE", "KODE CABANG", "NAMA CABANG", COL_KONTRAK, "NAMA DEBITUR",
    "NO AKTA", "TANGGAL AKTA", "NO SERTIFIKAT", "TANGGAL SERTIFIKAT",
    "BIAYA JASA", "PNBP", "TOTAL", COL_TGL_BAYAR,
]
DATE_HEADERS = {"TANGGAL INVOICE", "TANGGAL AKTA", "TANGGAL SERTIFIKAT", COL_TGL_BAYAR}
MONEY_HEADERS = {"BIAYA JASA", "PNBP", "TOTAL"}
DATE_FMT = "DD-MM-YYYY"


def _money(v):
    try:
        return Decimal(str(v).replace(",", "").strip()) if v not in (None, "") else Decimal("0")
    except InvalidOperation:
        return Decimal("0")


def outstanding_qs():
    return tb_invoice.objects.filter(tgl_receive_payment__isnull=True)


def history_qs():
    return tb_invoice.objects.filter(tgl_receive_payment__isnull=False)


def scope_to_branches(qs, branch_codes):
    """branch_codes None = tanpa batasan (user pusat)."""
    return qs if branch_codes is None else qs.filter(branch_id__in=branch_codes)


def billing_rows(invoice_qs):
    """Gabungkan tb_invoice dengan data konsumen & akta. Return list dict siap tampil/export."""
    invoices = list(invoice_qs.values(
        "id", "no_invoice", "tgl_inv", "branch_id", "agreement_number", "customer_id",
        "service_cost", "pnbp_fees", "tgl_receive_payment", "log_accept",
    ))
    consumer_ids = {i["customer_id"] for i in invoices}
    consumers = {
        c["id"]: c for c in Consumer.objects.filter(id__in=consumer_ids)
        .values("id", "debtor_name", "branch_name", "branch__branch_name")
    }
    docs = {
        d["consumer_id"]: d for d in LegalDocuments.objects.filter(consumer_id__in=consumer_ids)
        .values("consumer_id", "document_number", "document_date", "certificate_number", "certificate_date")
    }

    rows = []
    for inv in invoices:
        cons = consumers.get(inv["customer_id"], {})
        doc = docs.get(inv["customer_id"], {})
        service, pnbp = _money(inv["service_cost"]), _money(inv["pnbp_fees"])
        rows.append({
            "id": inv["id"],
            "no_invoice": inv["no_invoice"],
            "tgl_inv": inv["tgl_inv"],
            "branch_code": inv["branch_id"] or "",
            "branch_name": cons.get("branch__branch_name") or cons.get("branch_name") or "",
            "agreement_number": inv["agreement_number"],
            "debtor_name": cons.get("debtor_name") or "",
            "document_number": doc.get("document_number") or "",
            "document_date": doc.get("document_date"),
            "certificate_number": doc.get("certificate_number") or "",
            "certificate_date": doc.get("certificate_date"),
            "service_fee": service,
            "pnbp_fee": pnbp,
            "total": service + pnbp,
            "tgl_receive_payment": inv["tgl_receive_payment"],
            "paid_by": (inv["log_accept"] or "").split("|")[0],
        })
    return rows


def _export_values(r):
    return [
        r["no_invoice"], r["tgl_inv"], r["branch_code"], r["branch_name"], r["agreement_number"],
        r["debtor_name"], str(r["document_number"]), r["document_date"], r["certificate_number"],
        r["certificate_date"], r["service_fee"], r["pnbp_fee"], r["total"], r["tgl_receive_payment"],
    ]


def build_export_workbook(rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(_export_values(r))
    for col_idx, header in enumerate(EXPORT_HEADERS, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = max(14, len(header) + 2)
        for (cell,) in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            if header in DATE_HEADERS:
                cell.number_format = DATE_FMT
            elif header in MONEY_HEADERS:
                cell.number_format = "#,##0"
            else:
                cell.number_format = "@"  # jaga nol di depan no. kontrak / no. akta
    ws.freeze_panes = "A2"
    return wb


# ---------------------------------------------------------------------------
# Upload pembayaran
# ---------------------------------------------------------------------------

def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_payment_date(v):
    """Tanggal bayar dari sel Excel (tanggal asli atau teks DD-MM-YYYY / DD/MM/YYYY / YYYY-MM-DD)."""
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(s)


def parse_payment_rows(rows, today=None):
    """
    rows: tuple per baris dari sheet (baris pertama = header).
    Return (entries, unpaid_rows, errors). entries = [(baris_excel, no_invoice, no_kontrak, tgl_bayar)].
    """
    today = today or timezone.localdate()
    if not rows:
        return [], [], ["File kosong."]
    header = [_cell_str(h).upper() for h in rows[0]]
    missing = [c for c in (COL_INVOICE, COL_KONTRAK, COL_TGL_BAYAR) if c not in header]
    if missing:
        return [], [], [f"Kolom wajib tidak ditemukan: {', '.join(missing)}. Gunakan file hasil Export Billing."]
    i_inv, i_kon, i_tgl = (header.index(c) for c in (COL_INVOICE, COL_KONTRAK, COL_TGL_BAYAR))

    entries, unpaid, errors = [], [], []
    for excel_row, row in enumerate(rows[1:], start=2):
        cells = list(row) + [None] * len(header)
        no_inv, no_kon = _cell_str(cells[i_inv]), _cell_str(cells[i_kon])
        raw_tgl = cells[i_tgl]
        if not no_inv and not no_kon and raw_tgl in (None, ""):
            continue
        if not no_inv or not no_kon:
            errors.append(f"Baris {excel_row}: NOMOR INVOICE dan NOMOR KONTRAK wajib diisi.")
            continue
        try:
            tgl = parse_payment_date(raw_tgl)
        except ValueError:
            errors.append(f"Baris {excel_row}: TANGGAL BAYAR '{_cell_str(raw_tgl)}' tidak dikenali (gunakan DD-MM-YYYY).")
            continue
        if tgl is None:
            unpaid.append(excel_row)
        elif tgl > today:
            errors.append(f"Baris {excel_row}: TANGGAL BAYAR {tgl:%d-%m-%Y} melebihi hari ini.")
        else:
            entries.append((excel_row, no_inv, no_kon, tgl))
    return entries, unpaid, errors


def apply_payments(entries, user, branch_codes=None):
    """Tandai kontrak lunas. Hanya baris outstanding yang diubah; yang sudah dibayar tidak ditimpa."""
    result = {"paid": [], "already_paid": [], "not_found": []}
    stamp = f"{getattr(user, 'username', '')}|{timezone.now():%Y-%m-%d %H:%M:%S}"
    with transaction.atomic():
        for excel_row, no_inv, no_kon, tgl in entries:
            qs = scope_to_branches(
                tb_invoice.objects.select_for_update().filter(no_invoice=no_inv, agreement_number=no_kon),
                branch_codes,
            )
            item = {"row": excel_row, "no_invoice": no_inv, "agreement_number": no_kon, "tgl_bayar": tgl.isoformat()}
            invoices = list(qs)
            if not invoices:
                result["not_found"].append(item)
                continue
            open_ids = [inv.id for inv in invoices if inv.tgl_receive_payment is None]
            if not open_ids:
                result["already_paid"].append(item)
                continue
            tb_invoice.objects.filter(id__in=open_ids).update(
                tgl_receive_payment=tgl, log_accept=stamp, updated_at=timezone.now()
            )
            result["paid"].append(item)
    return result
