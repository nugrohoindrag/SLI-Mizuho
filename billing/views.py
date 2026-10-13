import os
import re
from datetime import timedelta

from django.contrib.auth.decorators import login_required
from django.core.files.storage import default_storage
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST
from openpyxl import load_workbook

from AdminMenu.auth_helpers import get_user_permissions, permission_required
from AdminUser.branch_access import get_user_branch_codes, user_is_branch_restricted
from common.validators import validate_excel_file

from . import services
from .models import BillingPaymentUpload

UPLOAD_DIR = "billing-upload"
MAX_RANGE_DAYS = 92


def _branch_codes(user):
    """None = semua cabang; list = dibatasi cabang user."""
    if not user_is_branch_restricted(user):
        return None
    return list(get_user_branch_codes(user) or [])


def _can_upload(user):
    return user.is_superuser or "upload-billing" in get_user_permissions(user)


def _date_range(request, from_key, to_key):
    """Return (start, end, error). Kosong keduanya = tanpa filter tanggal."""
    raw_from, raw_to = request.GET.get(from_key, ""), request.GET.get(to_key, "")
    if not raw_from and not raw_to:
        return None, None, None
    start, end = parse_date(raw_from or raw_to), parse_date(raw_to or raw_from)
    if not start or not end:
        return None, None, "Format tanggal tidak valid."
    if start > end:
        return None, None, "Tanggal awal tidak boleh lebih besar dari tanggal akhir."
    if (end - start).days > MAX_RANGE_DAYS:
        return None, None, f"Rentang tanggal maksimal {MAX_RANGE_DAYS} hari."
    return start, end, None


def _outstanding(request):
    start, end, error = _date_range(request, "tgl_inv_from", "tgl_inv_to")
    qs = services.scope_to_branches(services.outstanding_qs(), _branch_codes(request.user))
    if error:
        return qs.none(), start, end, error
    if start:
        qs = qs.filter(tgl_inv__range=(start, end))
    return qs.order_by("tgl_inv", "no_invoice", "agreement_number"), start, end, None


def _history(request):
    q = (request.GET.get("q") or "").strip()
    start, end, error = _date_range(request, "tgl_bayar_from", "tgl_bayar_to")
    qs = services.scope_to_branches(services.history_qs(), _branch_codes(request.user))
    if error:
        return qs.none(), q, start, end, error
    if q:
        qs = qs.filter(Q(no_invoice__icontains=q) | Q(agreement_number__icontains=q))
    if start:
        qs = qs.filter(tgl_receive_payment__range=(start, end))
    elif not q:
        # Tanpa pencarian: default 30 hari pembayaran terakhir agar tabel tidak membengkak
        end = timezone.localdate()
        start = end - timedelta(days=30)
        qs = qs.filter(tgl_receive_payment__range=(start, end))
    return qs.order_by("-tgl_receive_payment", "no_invoice", "agreement_number"), q, start, end, None


def _totals(rows):
    return {
        "count": len(rows),
        "service_fee": sum(r["service_fee"] for r in rows),
        "pnbp_fee": sum(r["pnbp_fee"] for r in rows),
        "total": sum(r["total"] for r in rows),
        "invoices": len({r["no_invoice"] for r in rows}),
    }


def _xlsx_response(wb, filename):
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required(login_url="")
@permission_required("read-billing")
def billing_list(request):
    qs, start, end, error = _outstanding(request)
    rows = services.billing_rows(qs)
    return render(request, "billing/billing.html", {
        "rows": rows,
        "totals": _totals(rows),
        "tgl_inv_from": start.isoformat() if start else "",
        "tgl_inv_to": end.isoformat() if end else "",
        "error_message": error or "",
        "can_upload": _can_upload(request.user),
    })


@login_required(login_url="")
@permission_required("read-billing")
def billing_export(request):
    qs, start, end, error = _outstanding(request)
    if error:
        return HttpResponse(error, status=400)
    suffix = f"{start:%Y%m%d}_{end:%Y%m%d}" if start else timezone.localdate().strftime("%Y%m%d")
    wb = services.build_export_workbook(services.billing_rows(qs), "Billing Outstanding")
    return _xlsx_response(wb, f"billing_outstanding_{suffix}.xlsx")


@login_required(login_url="")
@permission_required("upload-billing")
@require_POST
def billing_upload(request):
    uploaded = request.FILES.get("excel_file")
    if not uploaded:
        return JsonResponse({"status": "error", "message": "File belum dipilih."}, status=400)
    ok, message = validate_excel_file(uploaded)
    if not ok:
        return JsonResponse({"status": "error", "message": message}, status=400)

    try:
        wb = load_workbook(uploaded, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    except Exception:
        return JsonResponse({"status": "error", "message": "File Excel tidak dapat dibaca."}, status=400)

    entries, unpaid, errors = services.parse_payment_rows(rows)
    if errors:
        # Semua-atau-tidak sama sekali: perbaiki file dulu agar tidak ada pembayaran setengah jalan
        return JsonResponse({"status": "error", "message": "Cek kembali Excel Anda.", "errors": errors}, status=400)
    if not entries:
        return JsonResponse({"status": "error", "message": "Tidak ada baris dengan TANGGAL BAYAR terisi."}, status=400)

    result = services.apply_payments(entries, request.user, _branch_codes(request.user))

    uploaded.seek(0)
    safe_name = re.sub(r"[^0-9A-Za-z._-]", "_", os.path.basename(uploaded.name))
    path = default_storage.save(f"{UPLOAD_DIR}/{timezone.now():%Y%m%d%H%M%S}_{safe_name}", uploaded)
    BillingPaymentUpload.objects.create(
        file_name=uploaded.name[:255],
        file_path=path,
        uploaded_by=request.user,
        total_rows=len(entries) + len(unpaid),
        paid_count=len(result["paid"]),
        unpaid_count=len(unpaid),
        already_paid_count=len(result["already_paid"]),
        not_found_count=len(result["not_found"]),
        invalid_count=0,
        details=result,
    )

    return JsonResponse({
        "status": "success",
        "summary": {
            "paid": len(result["paid"]),
            "unpaid": len(unpaid),
            "already_paid": len(result["already_paid"]),
            "not_found": len(result["not_found"]),
        },
        "already_paid": result["already_paid"][:50],
        "not_found": result["not_found"][:50],
    })


@login_required(login_url="")
@permission_required("read-billing")
def billing_history(request):
    qs, q, start, end, error = _history(request)
    rows = services.billing_rows(qs)
    uploads = BillingPaymentUpload.objects.select_related("uploaded_by")[:20]
    return render(request, "billing/history.html", {
        "rows": rows,
        "totals": _totals(rows),
        "q": q,
        "tgl_bayar_from": start.isoformat() if start else "",
        "tgl_bayar_to": end.isoformat() if end else "",
        "error_message": error or "",
        "uploads": uploads,
    })


@login_required(login_url="")
@permission_required("read-billing")
def billing_history_export(request):
    qs, q, start, end, error = _history(request)
    if error:
        return HttpResponse(error, status=400)
    wb = services.build_export_workbook(services.billing_rows(qs), "History Pembayaran")
    return _xlsx_response(wb, f"billing_history_{timezone.localdate():%Y%m%d}.xlsx")
