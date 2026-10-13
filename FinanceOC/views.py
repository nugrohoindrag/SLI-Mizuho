from django.shortcuts import render, get_object_or_404, redirect
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from import_cabang_oc.models import ConsumerOC, Branch
from warkahDokumentOC.models import LegalDocumentOC
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from openpyxl import Workbook, load_workbook
from django.utils.dateparse import parse_date
from AktaSertifikat.models import Notaries
from AktaSertifikat.views import determine_status, can_view_all
from AdminUser.branch_access import get_user_branch_codes, user_is_branch_restricted
from AdminMenu.auth_helpers import permission_required, get_user_role_names
from common.const.role_map import has_role, ROLE_KEYS
from common.validators import validate_excel_file_or_error
from django.db.models import F, CharField, Q
from django.db.models.functions import Coalesce
from datetime import date, datetime
from io import BytesIO

# Create your views here.

def _user_is_staff_cabang(user):
    """Check if user has staff_cabang role"""
    return user_is_branch_restricted(user)

def _filter_docs_by_branch(qs, user):
    if not _user_is_staff_cabang(user):
        return qs

    branch_codes = get_user_branch_codes(user)
    if not branch_codes:
        return qs.none()

    return qs.filter(
        Q(consumer__branch__branch_code__in=branch_codes) | Q(consumer__branch_id__in=branch_codes)
    )

@login_required(login_url="")
@permission_required("read-finance-reg")
def finance_list_oc(request):
    document_date_str = request.GET.get("deed_date")
    
    if not document_date_str:
        document_date_str = date.today().strftime('%Y-%m-%d')

    legal_docs_qs = (
        LegalDocumentOC.objects
        .filter(document_date=document_date_str)
        .order_by("-created_at")
    )
    legal_docs_qs = _filter_docs_by_branch(legal_docs_qs, request.user)
    
    docs_data = list(
        legal_docs_qs.values(
            "document_number", "document_date", "voucher_code", 
            "certificate_number", "certificate_date", "payment_date",
            "registration_date", "certificate_path",
            
            agreement_number=F('consumer__agreement_number'),
            debtor_name=F('consumer__debtor_name'),
            guarantee_value=F('consumer__guarantee_value'),
            fee_pnbp=F('consumer__pnbp_fee'),
            fee_service=F('consumer__service_fee'),
            notary_name=F('notary__firm_name'),
            branch_code=F('consumer__branch__branch_code'),
            branch_name=Coalesce(
                F('consumer__branch__branch_name'),
                F('consumer__branch_name'),
                output_field=CharField()
            )
        )
    )
    
    for doc_item in docs_data:
        doc_item['status'] = determine_status(doc_item)

    context = {
        "enriched_docs": docs_data,
        "deed_date": document_date_str,
    }
    return render(request, "Finance/page-finance-OC.html", context)


@login_required(login_url="")
@permission_required("read-finance-reg")
def export_excel_finance_oc(request):
    export_status = (request.GET.get("export_status") or "").lower()
    deed_date = request.GET.get("deed_date")

    if export_status == "check" and not deed_date:
        return HttpResponse("Tanggal akta harus diisi untuk status CHECK.", status=400)

    legal_docs = LegalDocumentOC.objects.select_related(
        "consumer", "consumer__branch", "notary"
    )

    if deed_date:
        legal_docs = legal_docs.filter(document_date=deed_date)

    legal_docs = _filter_docs_by_branch(legal_docs, request.user)

    legal_docs = legal_docs.order_by("id")

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Data Sertifikat")

    ws.append(
        [
            "ID Konsumen",
            "Kode Cabang",
            "Nama Cabang",
            "No. Akta",
            "Tanggal Akta",
            "No. Perjanjian",
            "Nama Debitur",
            "Kode Voucher",
            "No. Sertifikat",
            "Tanggal Sertifikat",
            "Nilai Penjaminan",
            "Nominal PNBP",
            "Jasa",
            "Nama Notaris",
            "Tanggal Pembayaran",
            "Status",
        ]
    )

    for doc in legal_docs.iterator(chunk_size=1000):
        doc_item = {
            "voucher_code": doc.voucher_code,
            "registration_date": doc.registration_date,
            "certificate_number": doc.certificate_number,
            "payment_date": doc.payment_date,
            "certificate_path": doc.certificate_path,
            "certificate_date": doc.certificate_date,
        }

        status = determine_status(doc_item).lower()

        if export_status and status != export_status:
            continue

        consumer = doc.consumer

        branch_code = ""
        branch_name = ""

        if consumer:
            if consumer.branch:
                branch_code = consumer.branch.branch_code or ""
                branch_name = consumer.branch.branch_name or ""
            else:
                branch_name = consumer.branch_name or ""

        ws.append(
            [
                consumer.id if consumer else "",
                branch_code,
                branch_name,
                doc.document_number or "",
                doc.document_date,
                consumer.agreement_number if consumer else "",
                consumer.debtor_name if consumer else "",
                doc.voucher_code or "",
                doc.certificate_number or "",
                doc.certificate_date,
                consumer.guarantee_value if consumer else "",
                consumer.pnbp_fee if consumer else "",
                consumer.service_fee if consumer else "",
                doc.notary.firm_name if doc.notary else "",
                doc.payment_date,
                status.upper(),
            ]
        )

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    if export_status == "check" and deed_date:
        filename = f"sertifikat_{deed_date}_status_check_oc.xlsx"
    else:
        filename = f"sertifikat_status_clear_oc.xlsx"

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required(login_url="")
@permission_required("read-finance-reg")
def import_excel_finance_oc(request):
    if request.method == "POST":
        branch_codes = []
        if _user_is_staff_cabang(request.user):
            branch_codes = get_user_branch_codes(request.user)
            if not branch_codes:
                return JsonResponse({"message": "Forbidden"}, status=403)

        file = request.FILES.get("excel_file")
        if not file:
            return JsonResponse({"message": "File tidak ditemukan"}, status=400)

        # Validate Excel file
        validation_error = validate_excel_file_or_error(file)
        if validation_error:
            return validation_error

        wb = load_workbook(filename=file, read_only=True)
        ws = wb.active

        updated_count = 0
        skipped_details = []
        failed_details = []

        CERT_COL_INDEX = 8
        PAYMENT_DATE_COL_INDEX = 14

        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                consumer_id = row[0].value
                if not consumer_id:
                    continue

                cert_number_cell_value = row[CERT_COL_INDEX].value
                cert_number_val_upper = str(cert_number_cell_value or '').strip().upper()

                if "CLEAR" in cert_number_val_upper or "GAGAL" in cert_number_val_upper:
                    consumer_qs = ConsumerOC.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        failed_details.append(f"Baris {idx}: ConsumerOC ID '{consumer_id}' tidak ditemukan.")
                        continue

                    doc = LegalDocumentOC.objects.filter(consumer=consumer).first()
                    if not doc:
                        failed_details.append(f"Baris {idx}: Dokumen legal untuk ConsumerOC ID '{consumer_id}' tidak ditemukan.")
                        continue

                    doc.certificate_number = cert_number_cell_value
                    doc.payment_date = parse_date(row[PAYMENT_DATE_COL_INDEX].value)
                    doc.save(update_fields=['certificate_number', 'payment_date'])
                    
                    updated_count += 1
                else:
                    reason = f"Status '{cert_number_cell_value}' tidak diizinkan" if cert_number_cell_value else "Status kosong"
                    skipped_details.append(f"Baris {idx}: Dilewati ({reason})")
                    continue

            except Exception as e:
                failed_details.append(f"Baris {idx}: Terjadi error ({str(e)})")
                continue
        
        response_data = {
            "updated_count": updated_count,
            "skipped_count": len(skipped_details),
            "failed_count": len(failed_details),
            "skipped_details": skipped_details[:5],
            "failed_details": failed_details[:5],
        }
        return JsonResponse(response_data)

    return JsonResponse({"message": "Metode request tidak valid"}, status=400)
