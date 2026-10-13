import html
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, get_object_or_404, redirect
from import_cabang.models import Consumer, ChangeOfFiduciary
from import_cabang_oc.models import ConsumerOC
from import_cabang_mu.models import ConsumerMU
from django.core.paginator import Paginator
from django.utils.dateparse import parse_date
from django.utils import timezone
from django.db.models import (
    Q,
    Prefetch,
    OuterRef,
    Subquery,
    Window,
    F,
    When,
    Case,
    Value,
    IntegerField,
    Max,
    CharField,
)
from django.db.models.functions import RowNumber, Coalesce, Cast
from django.views.decorators.csrf import csrf_exempt, csrf_protect
from django.http import (
    FileResponse,
    Http404,
    HttpResponse,
    HttpResponseRedirect,
    JsonResponse,
    HttpResponseForbidden,
)
from import_document_mu.models import ConsumerDocumentMU
from .models import (
    Notaries,
    LegalDocuments,
    AktaMinutaLog,
    AktaMinutaLogOC,
    AktaMinutaLogMU,
    AktaMinutaLogPF,
    LegalDocumentsPF,
    LegalDocumentsOC,
    CertificateDownloadLogs,
    DeedDownloadLogs,
)
from warkahDokumentMU.models import LegalDocumentMU
import pandas as pd
from warkahDokumentOC.models import LegalDocumentOC
from common.validators import validate_excel_file_or_error
import mimetypes, os, json
from openpyxl import Workbook, load_workbook
from datetime import datetime, time, timedelta, date
from AdminUser.utils import check_permission
from AdminMenu.auth_helpers import (
    permission_required,
    get_user_role_names,
    user_has_permission,
)
from common.const.role_map import has_role, ROLE_KEYS
from AdminUser.branch_access import (
    get_primary_user_branch_code,
    get_user_branch_codes as _get_user_branch_codes,
    user_is_branch_restricted,
)
from django.urls import reverse
from django.core.exceptions import FieldError
from urllib.parse import quote
from io import BytesIO
from sysconfig import get_path
from django.core.files.storage import default_storage


def _user_is_superadmin(user):
    role_names = list(get_user_role_names(user))
    return (
        has_role(role_names, ROLE_KEYS["SUPERADMIN"])
        or has_role(role_names, ROLE_KEYS["HO"])
        or has_role(role_names, ROLE_KEYS["OPERATOR"])
        or has_role(role_names, ROLE_KEYS["PRINTING"])
        or has_role(role_names, ROLE_KEYS["WARKAH"])
        or getattr(user, "is_superuser", False)
    )


def can_view_all(user):
    """
    True jika user boleh melihat semua cabang.
    Atur sesuai role/flag yang Anda pakai.
    """
    return _user_is_superadmin(user)


def _user_is_staff_cabang(user):
    """Check if user has staff_cabang role"""
    return user_is_branch_restricted(user)


def get_user_branch_code(user):
    """
    Ambil branch_code milik user.
    - Jika FK user.branch ke Branch(pk=branch_code), maka user.branch_id sudah berisi branch_code (string).
    """
    return get_primary_user_branch_code(user)


def get_user_branch_codes(user):
    return _get_user_branch_codes(user)


def filter_queryset_by_user_branch(qs, user):
    """
    Filter queryset berdasarkan branch user.
    Diasumsikan model punya salah satu dari:
    - field 'branch_code' (CharField)
    - FK 'branch' ke model Branch yang PK-nya = branch_code
    Jika user bukan staff_cabang → kembalikan qs tanpa filter.
    Jika user tidak punya branch → kembalikan qs.none() untuk aman.
    """
    if not _user_is_staff_cabang(user):
        return qs

    branch_codes = get_user_branch_codes(user)
    if not branch_codes:
        return qs.none()

    # Coba filter dengan beberapa kemungkinan nama field
    for field_name in ("branch_code", "branch_id", "branch__branch_code", "branch__pk"):
        try:
            return qs.filter(**{f"{field_name}__in": branch_codes})
        except FieldError:
            continue

    # Jika semua gagal, aman-ketat
    return qs.none()


def _filter_by_branch(qs, user, branch_fields):
    if not _user_is_staff_cabang(user):
        return qs

    branch_codes = get_user_branch_codes(user)
    if not branch_codes:
        return qs.none()

    for field_name in branch_fields:
        try:
            return qs.filter(**{f"{field_name}__in": branch_codes})
        except FieldError:
            continue

    return qs.none()


def _write_branch_codes_or_forbidden(user):
    if not _user_is_staff_cabang(user):
        return []

    branch_codes = get_user_branch_codes(user)
    return branch_codes or None


def _can_export_akser_pf(user):
    return (
        getattr(user, "is_superuser", False)
        or getattr(user, "username", "").lower() == "admin"
        or user_has_permission(user, "read-akser-pf")
        or _user_is_staff_cabang(user)
    )


def _agreement_numbers_for_branch(model, branch_code):
    branch_codes = [branch_code] if branch_code else []
    return _agreement_numbers_for_branches(model, branch_codes)


def _agreement_numbers_for_branches(model, branch_codes):
    if not branch_codes:
        return set()

    base_qs = model.objects.filter(
        Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
    )
    agreements = set(base_qs.values_list("agreement_number", flat=True))

    try:
        model._meta.get_field("new_agreement_number")
        agreements.update(base_qs.values_list("new_agreement_number", flat=True))
    except Exception:
        pass

    return {value for value in agreements if value}


def _get_branch_code_from_consumer(consumer):
    if not consumer:
        return None
    if getattr(consumer, "branch", None):
        return getattr(consumer.branch, "branch_code", None) or getattr(
            consumer.branch, "pk", None
        )
    return getattr(consumer, "branch_code", None) or getattr(consumer, "branch_id", None)


def _user_can_access_branch(user, branch_code):
    if not _user_is_staff_cabang(user):
        return True
    user_branches = get_user_branch_codes(user)
    if not user_branches or not branch_code:
        return False
    return str(branch_code) in {str(code) for code in user_branches}


def _deed_date_consumer_ids(start_date, end_date):
    """Consumer yang aktanya bertanggal (legal_documents.document_date) dalam rentang."""
    return LegalDocuments.objects.filter(
        document_date__range=(start_date, end_date)
    ).values("consumer_id")


@login_required(login_url="")
@permission_required("read-akser")
def akta_sertifikat_list_view(request):
    consumers_qs = Consumer.objects.all().order_by("-created_at")
    consumers_qs = _filter_by_branch(
        consumers_qs,
        request.user,
        ["branch__branch_code", "branch_id", "branch__pk"],
    )

    tanggal_mulai = request.GET.get("tanggal_mulai")
    tanggal_selesai = request.GET.get("tanggal_selesai")

    error_message = ""
    today = date.today()

    if tanggal_mulai and tanggal_selesai:
        start_date = parse_date(tanggal_mulai)
        end_date = parse_date(tanggal_selesai)

        if not start_date or not end_date:
            error_message = "Format tanggal tidak valid"
        elif start_date > end_date:
            error_message = "Tanggal mulai tidak boleh lebih besar dari tanggal selesai"
        elif (end_date - start_date).days > 30:
            error_message = "Rentang tanggal maksimal 30 hari"
        else:
            consumers_qs = consumers_qs.filter(
                id__in=_deed_date_consumer_ids(start_date, end_date)
            )
        if error_message:
            consumers_qs = consumers_qs.none()
    else:
        consumers_qs = consumers_qs.filter(id__in=_deed_date_consumer_ids(today, today))
        tanggal_mulai = today.strftime("%Y-%m-%d")
        tanggal_selesai = today.strftime("%Y-%m-%d")

    consumer_fields = [
        "id",
        "fiducia_provider",
        "debtor_name",
        "name_on_bpkb",
        "agreement_number",
        "branch_name",
        "pnbp_fee",
        "document_status",
        "created_at",
    ]
    consumers_data = list(consumers_qs.values(*consumer_fields))

    consumer_ids = [c["id"] for c in consumers_data]

    legal_docs_qs = LegalDocuments.objects.filter(
        consumer_id__in=consumer_ids
    ).select_related("notary")
    document_map = {doc.consumer_id: doc for doc in legal_docs_qs}
    
    from import_document.models import ConsumerDocument
    agreement_numbers = [c["agreement_number"] for c in consumers_data if c["agreement_number"]]
    consumer_docs_qs = (
        ConsumerDocument.objects
        .filter(agreement_number__in=agreement_numbers)
        .values("agreement_number")
        .annotate(doc_created_at=Max("created_at"))
    )
    consumer_doc_map = {d["agreement_number"]: d["doc_created_at"] for d in consumer_docs_qs}

    status_labels = {"1": "Completed", "2": "Process", "3": "Revision"}

    for consumer_data in consumers_data:
        doc = document_map.get(consumer_data["id"])

        consumer_data["document_status_label"] = status_labels.get(
            consumer_data["document_status"], "Uncompleted"
        )

        is_process_done = consumer_data["document_status"] in ["1", "2", "3"]
        consumer_data["is_process_done"] = is_process_done

        consumer_data["has_legal_document"] = doc is not None

        consumer_data["can_download_akta"] = (
            is_process_done and doc and bool(doc.pnbp_path)
        )
        consumer_data["can_download_certificate"] = (
            is_process_done and doc and bool(doc.certificate_path)
        )
        consumer_data["can_download_bukti_bayar"] = bool(doc and doc.voucher_input)

        provider = (consumer_data["fiducia_provider"] or "").lower()
        if "atas nama customer" in provider:
            display_fiducia = consumer_data["debtor_name"]
        elif "atas nama bpkb" in provider:
            display_fiducia = consumer_data["name_on_bpkb"]
        else:
            display_fiducia = consumer_data["fiducia_provider"]

        consumer_data["display_fiducia"] = display_fiducia
        consumer_data["firm_name"] = doc.notary.firm_name if doc and doc.notary else ""
        consumer_data["voucher_code"] = doc.voucher_code if doc else ""
        consumer_data["certificate_number"] = doc.certificate_number if doc else ""
        consumer_data["certificate_date"] = doc.certificate_date if doc else ""
        consumer_data["document_id"] = doc.id if doc else None
        consumer_data["document_number"] = doc.document_number if doc else ""
        consumer_data["document_date"] = doc.document_date if doc else None
        consumer_data["consumer_document_created_at"] = consumer_doc_map.get(
            consumer_data["agreement_number"]
        )

    context = {
        "enriched_consumers": consumers_data,
        "tanggal_mulai": tanggal_mulai or "",
        "tanggal_selesai": tanggal_selesai or "",
        "error_message": error_message,
    }
    return render(request, "AktaSertifikat/consumer-list.html", context)


@login_required(login_url="")
@permission_required("read-sertifikat")
def download_sertifikat(request, id):
    document = get_object_or_404(LegalDocuments, pk=id)
    consumer = getattr(document, "consumer", None)
    if not _user_can_access_branch(
        request.user, _get_branch_code_from_consumer(consumer)
    ):
        return HttpResponseForbidden("Forbidden")

    file_rel_path = document.certificate_path
    if not file_rel_path:
        raise Http404("File sertifikat tidak tersedia")

    file_full_path = os.path.join(settings.MEDIA_ROOT, file_rel_path)
    print("Full file path:", file_full_path)

    if not os.path.exists(file_full_path):
        raise Http404("File sertifikat tidak ditemukan di server")

    try:
        CertificateDownloadLogs.objects.create(
            user_id=request.user.id,
            agreement_number=getattr(consumer, "agreement_number", "UNKNOWN"),
            created_at=timezone.now(),
            updated_at=timezone.now(),
        )
        print(f"✅ Log download tersimpan untuk user {request.user} dan agreement {consumer.agreement_number}")
    except Exception as e:
        print(f"⚠️ Gagal menyimpan log download sertifikat: {e}")
        
    return FileResponse(
        open(file_full_path, "rb"),
        as_attachment=True,
        filename=os.path.basename(file_full_path),
    )


@login_required(login_url="")
@permission_required("read-sertifikat")
def download_sertifikat_pf(request, id):
    document = get_object_or_404(LegalDocumentsPF, pk=id)
    consumer = getattr(document, "change_of_fiduciary", None)
    if not _user_can_access_branch(
        request.user, _get_branch_code_from_consumer(consumer)
    ):
        return HttpResponseForbidden("Forbidden")

    file_rel_path = document.certificate_path
    if not file_rel_path:
        raise Http404("File sertifikat tidak tersedia")

    file_full_path = os.path.join(settings.MEDIA_ROOT, file_rel_path)
    print("Full file path:", file_full_path)

    if not os.path.exists(file_full_path):
        raise Http404("File sertifikat tidak ditemukan di server")

    return FileResponse(
        open(file_full_path, "rb"),
        as_attachment=True,
        filename=os.path.basename(file_full_path),
    )


@login_required(login_url="")
@permission_required("read-sertifikat")
def download_sertifikat_oc(request, id):
    document = get_object_or_404(LegalDocumentsOC, pk=id)
    consumer = getattr(document, "consumer", None)
    if not _user_can_access_branch(
        request.user, _get_branch_code_from_consumer(consumer)
    ):
        return HttpResponseForbidden("Forbidden")

    file_rel_path = document.certificate_path
    if not file_rel_path:
        raise Http404("File sertifikat tidak tersedia")

    file_full_path = os.path.join(settings.MEDIA_ROOT, file_rel_path)
    print("Full file path:", file_full_path)

    if not os.path.exists(file_full_path):
        raise Http404("File sertifikat tidak ditemukan di server")

    return FileResponse(
        open(file_full_path, "rb"),
        as_attachment=True,
        filename=os.path.basename(file_full_path),
    )


@login_required(login_url="")
@permission_required("read-sertifikat")
def download_sertifikat_mu(request, id):
    document = get_object_or_404(LegalDocumentMU, pk=id)
    consumer = getattr(document, "consumer", None)
    if not _user_can_access_branch(
        request.user, _get_branch_code_from_consumer(consumer)
    ):
        return HttpResponseForbidden("Forbidden")

    file_rel_path = document.certificate_path
    if not file_rel_path:
        raise Http404("File sertifikat tidak tersedia")

    file_full_path = os.path.join(settings.MEDIA_ROOT, file_rel_path)
    print("Full file path:", file_full_path)

    if not os.path.exists(file_full_path):
        raise Http404("File sertifikat tidak ditemukan di server")

    return FileResponse(
        open(file_full_path, "rb"),
        as_attachment=True,
        filename=os.path.basename(file_full_path),
    )


@login_required(login_url="")
@permission_required("read-akser")
def download_akta(request, id):
    document = get_object_or_404(LegalDocuments, pk=id)
    consumer = getattr(document, "consumer", None)
    if not _user_can_access_branch(
        request.user, _get_branch_code_from_consumer(consumer)
    ):
        return HttpResponseForbidden("Forbidden")

    file_rel_path = document.pnbp_path
    if not file_rel_path:
        raise Http404("File akta tidak tersedia")

    file_full_path = os.path.join(settings.MEDIA_ROOT, file_rel_path)
    print("Full file path:", file_full_path)

    if not os.path.exists(file_full_path):
        raise Http404("File akta tidak ditemukan di server")

    try:
        DeedDownloadLogs.objects.create(
            user_id=request.user.id,
            agreement_number=getattr(consumer, "agreement_number", "UNKNOWN"),
            created_at=timezone.now(),
            updated_at=timezone.now(),
        )
        print(f"✅ Log download tersimpan untuk user {request.user} dan agreement {consumer.agreement_number}")
    except Exception as e:
        print(f"⚠️ Gagal menyimpan log download sertifikat: {e}")
        
    return FileResponse(
        open(file_full_path, "rb"),
        as_attachment=True,
        filename=os.path.basename(file_full_path),
    )


@login_required(login_url="")
@permission_required("read-akser-pf")
def download_akta_pf(request, id):
    document = get_object_or_404(LegalDocumentsPF, pk=id)
    consumer = getattr(document, "change_of_fiduciary", None)
    if not _user_can_access_branch(
        request.user, _get_branch_code_from_consumer(consumer)
    ):
        return HttpResponseForbidden("Forbidden")

    file_rel_path = document.pnbp_path
    if not file_rel_path:
        raise Http404("File akta tidak tersedia")

    file_full_path = os.path.join(settings.MEDIA_ROOT, file_rel_path)
    print("Full file path:", file_full_path)

    if not os.path.exists(file_full_path):
        raise Http404("File akta tidak ditemukan di server")

    return FileResponse(
        open(file_full_path, "rb"),
        as_attachment=True,
        filename=os.path.basename(file_full_path),
    )


@login_required(login_url="")
@permission_required("upload-akser-oc")
def download_akta_oc(request, id):
    document = get_object_or_404(LegalDocumentOC, pk=id)
    consumer = getattr(document, "consumer", None)
    if not _user_can_access_branch(
        request.user, _get_branch_code_from_consumer(consumer)
    ):
        return HttpResponseForbidden("Forbidden")

    file_rel_path = document.pnbp_path
    if not file_rel_path:
        raise Http404("File akta tidak tersedia")

    file_full_path = os.path.join(settings.MEDIA_ROOT, file_rel_path)
    print("Full file path:", file_full_path)

    if not os.path.exists(file_full_path):
        raise Http404("File akta tidak ditemukan di server")

    return FileResponse(
        open(file_full_path, "rb"),
        as_attachment=True,
        filename=os.path.basename(file_full_path),
    )


@login_required(login_url="")
@permission_required("read-akser-mu")
def download_akta_mu(request, id):
    document = get_object_or_404(LegalDocumentMU, pk=id)
    consumer = getattr(document, "consumer", None)
    if not _user_can_access_branch(
        request.user, _get_branch_code_from_consumer(consumer)
    ):
        return HttpResponseForbidden("Forbidden")

    file_rel_path = document.pnbp_path
    if not file_rel_path:
        raise Http404("File akta tidak tersedia")

    file_full_path = os.path.join(settings.MEDIA_ROOT, file_rel_path)
    print("Full file path:", file_full_path)

    if not os.path.exists(file_full_path):
        raise Http404("File akta tidak ditemukan di server")

    return FileResponse(
        open(file_full_path, "rb"),
        as_attachment=True,
        filename=os.path.basename(file_full_path),
    )


@login_required(login_url="")
@permission_required("read-update-sertifikat")
def update_sertifikat_list(request):
    deed_date = request.GET.get("deed_date")
    deed_date_to = request.GET.get("deed_date_to")

    docs_qs = LegalDocuments.objects.all()

    if not deed_date and not deed_date_to:
        today = date.today().strftime("%Y-%m-%d")
        deed_date = today
        deed_date_to = today
        docs_qs = docs_qs.filter(document_date=today)

    elif deed_date and deed_date_to:
        try:
            d1 = datetime.strptime(deed_date, "%Y-%m-%d")
            d2 = datetime.strptime(deed_date_to, "%Y-%m-%d")

            selisih = (d2 - d1).days

            if selisih < 0:
                return HttpResponse("Tanggal tidak valid (dari > sampai).", status=400)

            if selisih > 7:
                messages.error(request, "Range tanggal maksimal 7 hari.")
                return redirect(request.path)

            docs_qs = docs_qs.filter(document_date__range=[deed_date, deed_date_to])

        except ValueError:
            return HttpResponse("Format tanggal tidak valid.", status=400)

    elif deed_date:
        docs_qs = docs_qs.filter(document_date=deed_date)

    docs_qs = (
        docs_qs
        .select_related("consumer", "notary")
        .filter(consumer__agreement_number__isnull=False)
        .exclude(consumer__agreement_number="")
        .filter(document_number__isnull=False)
        .exclude(document_number=0)
        .filter(consumer__debtor_name__isnull=False)
        .exclude(consumer__debtor_name="")
        .filter(notary__isnull=False)
        .order_by("-consumer__created_at")
    )
    docs_qs = _filter_by_branch(
        docs_qs,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    docs_data = list(
        docs_qs.values(
            "id",
            "document_number",
            "voucher_code",
            "certificate_number",
            "certificate_date",
            "registration_date",
            "certificate_path",
            "pnbp_path",
            "voucher_input",
            "payment_date",
            "document_date",
            firm_name=F("notary__firm_name"),
            agreement_number=F("consumer__agreement_number"),
            debtor_name=F("consumer__debtor_name"),
        )
    )

    for doc_item in docs_data:
        doc_item["status"] = determine_status(doc_item)

    context = {
        "enriched_consumers": docs_data,
        "deed_date": deed_date,
        "deed_date_to": deed_date_to,
    }
    return render(request, "AktaSertifikat/update-sertifikat.html", context)


@login_required(login_url="")
@permission_required("read-update-sertifikat")
def update_sertifikat_list_pf(request):
    document_date = request.GET.get("deed_date")
    document_number = request.GET.get("deed_number")

    legal_docs_qs = LegalDocumentsPF.objects.all().select_related(
        "notary", "change_of_fiduciary"
    )

    if document_date:
        legal_docs_qs = legal_docs_qs.filter(document_date=document_date)
    else:
        today = date.today()
        legal_docs_qs = legal_docs_qs.filter(document_date=today)
        document_date = today.strftime("%Y-%m-%d")

    if document_number:
        legal_docs_qs = legal_docs_qs.filter(document_number=document_number)

    legal_docs_qs = _filter_by_branch(
        legal_docs_qs,
        request.user,
        [
            "change_of_fiduciary__branch__branch_code",
            "change_of_fiduciary__branch_id",
            "change_of_fiduciary__branch__pk",
        ],
    )

    docs_data = list(
        legal_docs_qs.values(
            "id",
            "document_number",
            "document_date",
            "document_time",
            "voucher_code",
            "certificate_number",
            "certificate_date",
            "registration_date",
            "certificate_path",
            "pnbp_path",
            "payment_date",
            firm_name=F("notary__firm_name"),
            agreement_number=Coalesce(
                F("change_of_fiduciary__new_agreement_number"),
                F("change_of_fiduciary__agreement_number"),
                output_field=CharField(),
            ),
            debtor_name=Coalesce(
                F("change_of_fiduciary__new_debtor_name"),
                F("change_of_fiduciary__debtor_name"),
                output_field=CharField(),
            ),
            name_on_bpkb=Coalesce(
                F("change_of_fiduciary__new_name_on_bpkb"),
                F("change_of_fiduciary__name_on_bpkb"),
                output_field=CharField(),
            ),
            fiducia_provider=F("change_of_fiduciary__fiducia_provider"),
        ).order_by("-change_of_fiduciary__created_at")
    )

    for item in docs_data:
        item["status"] = determine_status(item)

        provider = (item["fiducia_provider"] or "").lower()
        if "atas nama customer" in provider:
            item["display_fiducia"] = item["debtor_name"]
        elif "atas nama bpkb" in provider:
            item["display_fiducia"] = item["name_on_bpkb"]
        else:
            item["display_fiducia"] = item["fiducia_provider"]

    context = {
        "enriched_consumers_pf": docs_data,
        "deed_date": document_date or "",
        "deed_number": document_number or "",
    }
    return render(request, "AktaSertifikat/update-sertifikat-pf.html", context)


@login_required(login_url="")
@permission_required("read-update-sertifikat")
def update_sertifikat_list_oc(request):
    deed_date = request.GET.get("deed_date")
    deed_date_to = request.GET.get("deed_date_to")

    docs_qs = LegalDocumentsOC.objects.all()

    if not deed_date and not deed_date_to:
        today = date.today().strftime("%Y-%m-%d")
        deed_date = today
        deed_date_to = today
        docs_qs = docs_qs.filter(document_date=today)

    elif deed_date and deed_date_to:
        try:
            d1 = datetime.strptime(deed_date, "%Y-%m-%d")
            d2 = datetime.strptime(deed_date_to, "%Y-%m-%d")

            selisih = (d2 - d1).days

            if selisih < 0:
                return HttpResponse("Tanggal tidak valid (dari > sampai).", status=400)

            if selisih > 7:
                messages.error(request, "Range tanggal maksimal 7 hari.")
                return redirect(request.path)

            docs_qs = docs_qs.filter(document_date__range=[deed_date, deed_date_to])

        except ValueError:
            return HttpResponse("Format tanggal tidak valid.", status=400)

    elif deed_date:
        docs_qs = docs_qs.filter(document_date=deed_date)

    docs_qs = (
        docs_qs
        .select_related("consumer", "notary")
        .filter(consumer__agreement_number__isnull=False)
        .exclude(consumer__agreement_number="")
        .filter(document_number__isnull=False)
        .exclude(document_number=0)
        .filter(consumer__debtor_name__isnull=False)
        .exclude(consumer__debtor_name="")
        .filter(notary__isnull=False)
        .order_by("-consumer__created_at")
    )
    docs_qs = _filter_by_branch(
        docs_qs,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    docs_data = list(
        docs_qs.values(
            "id",
            "document_number",
            "voucher_code",
            "certificate_number",
            "certificate_date",
            "registration_date",
            "certificate_path",
            "pnbp_path",
            "payment_date",
            "document_date",
            firm_name=F("notary__firm_name"),
            agreement_number=F("consumer__agreement_number"),
            debtor_name=F("consumer__debtor_name"),
        )
    )

    for doc_item in docs_data:
        doc_item["status"] = determine_status(doc_item)

    context = {
        "enriched_consumers": docs_data,
        "deed_date": deed_date,
        "deed_date_to": deed_date_to,
    }
    return render(request, "AktaSertifikat/update-sertifikat-oc.html", context)


@login_required(login_url="")
@permission_required("read-update-sertifikat")
def update_sertifikat_list_mu(request):
    deed_date_raw = request.GET.get("deed_date")
    deed_date_to_raw = request.GET.get("deed_date_to")
    deed_number = request.GET.get("deed_number")

    deed_date = None
    deed_date_to = None

    try:
        if deed_date_raw:
            deed_date = datetime.strptime(deed_date_raw, "%Y-%m-%d").date()
        if deed_date_to_raw:
            deed_date_to = datetime.strptime(deed_date_to_raw, "%Y-%m-%d").date()
    except ValueError:
        return HttpResponse("Format tanggal tidak valid.", status=400)

    if not deed_date and not deed_date_to:
        deed_date = date.today()
        deed_date_to = date.today()
        deed_date_raw = deed_date.strftime("%Y-%m-%d")
        deed_date_to_raw = deed_date_to.strftime("%Y-%m-%d")

    if deed_date and deed_date_to:
        selisih = (deed_date_to - deed_date).days

        if selisih < 0:
            return HttpResponse("Tanggal tidak valid (dari > sampai).", status=400)

        if selisih > 7:
            messages.error(request, "Range tanggal maksimal 7 hari.")
            return redirect(request.path)

    consumers_qs = ConsumerMU.objects.order_by(
        "agreement_number", "-created_at"
    ).distinct("agreement_number")

    consumers_qs = _filter_by_branch(
        consumers_qs,
        request.user,
        ["branch__branch_code", "branch_id", "branch__pk"],
    )

    if deed_date and deed_date_to:
        consumers_qs = consumers_qs.filter(
            legal_documents_mu__document_date__range=[deed_date, deed_date_to]
        )
    elif deed_date:
        consumers_qs = consumers_qs.filter(
            legal_documents_mu__document_date=deed_date
        )

    if deed_number:
        consumers_qs = consumers_qs.filter(
            legal_documents_mu__document_number__icontains=deed_number
        )

    enriched_data = list(
        consumers_qs.values(
            "agreement_number",
            "debtor_name",
            idLegalDocsMu=F("legal_documents_mu__id"),
            document_number=F("legal_documents_mu__document_number"),
            voucher_code=F("legal_documents_mu__voucher_code"),
            certificate_number=F("legal_documents_mu__certificate_number"),
            certificate_date=F("legal_documents_mu__certificate_date"),
            registration_date=F("legal_documents_mu__registration_date"),
            certificate_path=F("legal_documents_mu__certificate_path"),
            pnbp_path=F("legal_documents_mu__pnbp_path"),
            payment_date=F("legal_documents_mu__payment_date"),
            firm_name=F("legal_documents_mu__notary__firm_name"),
        )
    )

    for item in enriched_data:
        item["deed_number"] = item["document_number"]
        item["status"] = determine_status(item)

    context = {
        "enriched_consumers": enriched_data,
        "deed_date": deed_date_raw or "",
        "deed_date_to": deed_date_to_raw or "",
        "deed_number": deed_number or "",
    }
    return render(request, "AktaSertifikat/update-sertifikat-mu.html", context)


@login_required(login_url="")
@permission_required("read-finance-reg")
def finance_reg_list(request):
    document_date_str = request.GET.get("deed_date")

    if not document_date_str:
        document_date_str = date.today().strftime("%Y-%m-%d")

    legal_docs_qs = LegalDocuments.objects.filter(
        document_date=document_date_str
    ).order_by("-created_at")
    legal_docs_qs = _filter_by_branch(
        legal_docs_qs,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    docs_data = list(
        legal_docs_qs.values(
            "document_number",
            "document_date",
            "voucher_code",
            "certificate_number",
            "certificate_date",
            "payment_date",
            "registration_date",
            "certificate_path",
            agreement_number=F("consumer__agreement_number"),
            debtor_name=F("consumer__debtor_name"),
            guarantee_value=F("consumer__guarantee_value"),
            fee_pnbp=F("consumer__pnbp_fee"),
            # fee_service=F("consumer__service_fee"),
            notary_name=F("notary__firm_name"),
            branch_code=F("consumer__branch__branch_code"),
            branch_name=Coalesce(
                F("consumer__branch__branch_name"),
                F("consumer__branch_name"),
                output_field=CharField(),
            ),
        )
    )

    for doc_item in docs_data:
        doc_item["status"] = determine_status(doc_item)

    context = {
        "enriched_docs": docs_data,
        "deed_date": document_date_str,
    }
    return render(request, "AktaSertifikat/finance-reg.html", context)


@login_required(login_url="")
@permission_required("read-akser-pf")
def akta_sertifikat_pf_list_view(request):
    cof_qs = ChangeOfFiduciary.objects.all().order_by("-created_at")
    cof_qs = _filter_by_branch(
        cof_qs,
        request.user,
        ["branch__branch_code", "branch_id", "branch__pk"],
    )

    tanggal_mulai = request.GET.get("tanggal_mulai")
    tanggal_selesai = request.GET.get("tanggal_selesai")

    error_message = ""
    today = date.today()
    
    if tanggal_mulai and tanggal_selesai:
        start_date = parse_date(tanggal_mulai)
        end_date = parse_date(tanggal_selesai)

        if not start_date or not end_date:
            error_message = "Format tanggal tidak valid"
        elif start_date > end_date:
            error_message = "Tanggal mulai tidak boleh lebih besar dari tanggal selesai"
        elif (end_date - start_date).days > 30:
            error_message = "Rentang tanggal maksimal 30 hari"
        else:
            cof_qs = cof_qs.filter(
                created_at__date__range=(start_date, end_date)
            )
        if error_message:
            cof_qs = cof_qs.none()
    else:
        cof_qs = cof_qs.filter(created_at__date=today)
        tanggal_mulai = today.strftime("%Y-%m-%d")
        tanggal_selesai = today.strftime("%Y-%m-%d")

    cof_fields = [
        "id",
        "new_debtor_name",
        "debtor_name",
        "new_name_on_bpkb",
        "name_on_bpkb",
        "fiducia_provider",
        "agreement_number",
        "branch_name",
        "pnbp_fee",
        "document_status",
        "created_at",
    ]
    cof_data = list(cof_qs.values(*cof_fields))

    cof_ids = [c["id"] for c in cof_data]

    legal_docs_pf = LegalDocumentsPF.objects.filter(
        change_of_fiduciary_id__in=cof_ids
    ).select_related("notary")
    document_map = {doc.change_of_fiduciary_id: doc for doc in legal_docs_pf}
    
    from import_document_pf.models import ConsumerDocumentPF
    agreement_numbers = [c["agreement_number"] for c in cof_data if c["agreement_number"]]
    consumer_docs_qs = (
        ConsumerDocumentPF.objects
        .filter(agreement_number__in=agreement_numbers)
        .values("agreement_number")
        .annotate(doc_created_at=Max("created_at"))
    )
    consumer_doc_map = {d["agreement_number"]: d["doc_created_at"] for d in consumer_docs_qs}

    status_labels = {"1": "Completed", "2": "Process", "3": "Revision"}

    for cof_item in cof_data:
        document = document_map.get(cof_item["id"])

        is_process_done = cof_item["document_status"] in ["1", "2", "3"]
        cof_item["document_status_label"] = status_labels.get(
            cof_item["document_status"], "Uncompleted"
        )
        cof_item["is_process_done"] = is_process_done
        cof_item["has_legal_document"] = document is not None
        cof_item["can_download_akta"] = (
            is_process_done and document and bool(document.pnbp_path)
        )
        cof_item["can_download_certificate"] = (
            is_process_done and document and bool(document.certificate_path)
        )

        debtor_name = cof_item["new_debtor_name"] or cof_item["debtor_name"]
        name_on_bpkb = cof_item["new_name_on_bpkb"] or cof_item["name_on_bpkb"]

        cof_item["debtor_name"] = debtor_name
        cof_item["name_on_bpkb"] = name_on_bpkb

        provider = (cof_item["fiducia_provider"] or "").lower()
        if "atas nama customer" in provider:
            display_fiducia = debtor_name
        elif "atas nama bpkb" in provider:
            display_fiducia = name_on_bpkb
        else:
            display_fiducia = cof_item["fiducia_provider"]

        cof_item["display_fiducia"] = display_fiducia
        cof_item["firm_name"] = (
            document.notary.firm_name if document and document.notary else ""
        )
        cof_item["voucher_code"] = document.voucher_code if document else ""
        cof_item["certificate_number"] = document.certificate_number if document else ""
        cof_item["certificate_date"] = document.certificate_date if document else ""
        cof_item["document_id"] = document.id if document else None
        cof_item["document_number"] = document.document_number if document else ""
        cof_item["document_date"] = document.document_date if document else None
        cof_item["consumer_document_created_at"] = consumer_doc_map.get(
            cof_item["agreement_number"]
        )

    context = {
        "enriched_data": cof_data,
        "tanggal_mulai": tanggal_mulai or "",
        "tanggal_selesai": tanggal_selesai or "",
        "error_message": error_message,
    }
    return render(request, "staff_cabang/page-staff-cabang-akser-pf.html", context)


@login_required(login_url="")
@permission_required("read-akser")
def akta_sertifikat_oc_list_view(request):
    consumers_qs = ConsumerOC.objects.all().order_by("-created_at")
    consumers_qs = _filter_by_branch(
        consumers_qs,
        request.user,
        ["branch__branch_code", "branch_id", "branch__pk"],
    )

    tanggal_mulai = request.GET.get("tanggal_mulai")
    tanggal_selesai = request.GET.get("tanggal_selesai")
    
    error_message = ""
    today = date.today()
    
    if tanggal_mulai and tanggal_selesai:
        start_date = parse_date(tanggal_mulai)
        end_date = parse_date(tanggal_selesai)

        if not start_date or not end_date:
            error_message = "Format tanggal tidak valid"
        elif start_date > end_date:
            error_message = "Tanggal mulai tidak boleh lebih besar dari tanggal selesai"
        elif (end_date - start_date).days > 30:
            error_message = "Rentang tanggal maksimal 30 hari"
        else:
            consumers_qs = consumers_qs.filter(
                created_at__date__range=(start_date, end_date)
            )
        if error_message:
            consumers_qs = consumers_qs.none()
    else:
        consumers_qs = consumers_qs.filter(created_at__date=today)
        tanggal_mulai = today.strftime("%Y-%m-%d")
        tanggal_selesai = today.strftime("%Y-%m-%d")

    consumer_fields = [
        "id",
        "fiducia_provider",
        "debtor_name",
        "name_on_bpkb",
        "agreement_number",
        "branch_name",
        "pnbp_fee",
        "document_status",
        "created_at",
    ]
    consumers_data = list(consumers_qs.values(*consumer_fields))

    consumer_ids = [c["id"] for c in consumers_data]

    legal_docs = LegalDocumentOC.objects.filter(
        consumer_id__in=consumer_ids
    ).select_related("notary")
    document_map = {doc.consumer_id: doc for doc in legal_docs}
    
    from import_document_oc.models import ConsumerDocumentOC
    agreement_numbers = [c["agreement_number"] for c in consumers_data if c["agreement_number"]]
    consumer_docs_qs = (
        ConsumerDocumentOC.objects
        .filter(agreement_number__in=agreement_numbers)
        .values("agreement_number")
        .annotate(doc_created_at=Max("created_at"))
    )
    consumer_doc_map = {d["agreement_number"]: d["doc_created_at"] for d in consumer_docs_qs}

    status_labels = {"1": "Completed", "2": "Process", "3": "Revision"}

    for consumer_data in consumers_data:
        document = document_map.get(consumer_data["id"])

        is_process_done = consumer_data["document_status"] in ["1", "2", "3"]
        consumer_data["document_status_label"] = status_labels.get(
            consumer_data["document_status"], "Uncompleted"
        )
        consumer_data["is_process_done"] = is_process_done
        consumer_data["has_legal_document"] = document is not None
        consumer_data["can_download_akta"] = (
            is_process_done and document and bool(document.pnbp_path)
        )
        consumer_data["can_download_certificate"] = (
            is_process_done and document and bool(document.certificate_path)
        )

        provider = (consumer_data["fiducia_provider"] or "").lower()
        if "atas nama customer" in provider:
            display_fiducia = consumer_data["debtor_name"]
        elif "atas nama bpkb" in provider:
            display_fiducia = consumer_data["name_on_bpkb"]
        else:
            display_fiducia = consumer_data["fiducia_provider"]

        consumer_data["display_fiducia"] = display_fiducia
        consumer_data["firm_name"] = (
            document.notary.firm_name if document and document.notary else ""
        )
        consumer_data["voucher_code"] = document.voucher_code if document else ""
        consumer_data["certificate_number"] = (
            document.certificate_number if document else ""
        )
        consumer_data["certificate_date"] = (
            document.certificate_date if document else ""
        )
        consumer_data["document_id"] = document.id if document else None
        consumer_data["document_number"] = document.document_number if document else ""
        consumer_data["document_date"] = document.document_date if document else None
        consumer_data["consumer_document_created_at"] = consumer_doc_map.get(
            consumer_data["agreement_number"]
        )

    can_upload = check_permission(request.user, "upload-akser-oc")

    context = {
        "enriched_consumers": consumers_data,
        "tanggal_mulai": tanggal_mulai or "",
        "tanggal_selesai": tanggal_selesai or "",
        "error_message": error_message,
        "can_upload": can_upload,
    }
    return render(request, "staff_cabang/page-staff-cabang-akser-oc.html", context)


@login_required(login_url="")
@permission_required("read-akser-mu")
def akta_sertifikat_mu_list_view(request):
    consumers_qs = ConsumerMU.objects.filter(is_st_used=1, st_process=1).order_by("-created_at")
    consumers_qs = _filter_by_branch(
        consumers_qs,
        request.user,
        ["branch__branch_code", "branch_id", "branch__pk"],
    )

    tanggal_mulai = request.GET.get("tanggal_mulai")
    tanggal_selesai = request.GET.get("tanggal_selesai")
    
    error_message = ""
    today = date.today()
    
    if tanggal_mulai and tanggal_selesai:
        start_date = parse_date(tanggal_mulai)
        end_date = parse_date(tanggal_selesai)

        if not start_date or not end_date:
            error_message = "Format tanggal tidak valid"
        elif start_date > end_date:
            error_message = "Tanggal mulai tidak boleh lebih besar dari tanggal selesai"
        elif (end_date - start_date).days > 30:
            error_message = "Rentang tanggal maksimal 30 hari"
        else:
            consumers_qs = consumers_qs.filter(
                created_at__date__range=(start_date, end_date)
            )
        if error_message:
            consumers_qs = consumers_qs.none()
    else:
        consumers_qs = consumers_qs.filter(created_at__date=today)
        tanggal_mulai = today.strftime("%Y-%m-%d")
        tanggal_selesai = today.strftime("%Y-%m-%d")

    consumer_fields = [
        "id",
        "fiducia_provider",
        "debtor_name",
        "name_on_bpkb",
        "agreement_number",
        "branch_name",
        "pnbp_fee",
        "document_status",
        "created_at",
    ]
    consumers_data = list(consumers_qs.values(*consumer_fields))

    consumer_ids = [c["id"] for c in consumers_data]

    legal_docs = LegalDocumentMU.objects.filter(
        consumer_id__in=consumer_ids
    ).select_related("notary")
    document_map = {doc.consumer_id: doc for doc in legal_docs}
    
    agreement_numbers = [c["agreement_number"] for c in consumers_data if c["agreement_number"]]
    consumer_docs_qs = (
        ConsumerDocumentMU.objects
        .filter(agreement_number__in=agreement_numbers)
        .values("agreement_number")
        .annotate(doc_created_at=Max("created_at"))
    )
    consumer_doc_map = {d["agreement_number"]: d["doc_created_at"] for d in consumer_docs_qs}

    status_labels = {"1": "Completed", "2": "Process", "3": "Revision"}

    for consumer_data in consumers_data:
        document = document_map.get(consumer_data["id"])

        is_process_done = consumer_data["document_status"] in ["1", "2", "3"]
        consumer_data["document_status_label"] = status_labels.get(
            consumer_data["document_status"], "Uncompleted"
        )
        consumer_data["is_process_done"] = is_process_done
        consumer_data["has_legal_document"] = document is not None
        consumer_data["can_download_akta"] = (
            is_process_done and document and bool(document.pnbp_path)
        )
        consumer_data["can_download_certificate"] = (
            is_process_done and document and bool(document.certificate_path)
        )

        provider = (consumer_data["fiducia_provider"] or "").lower()
        if "atas nama customer" in provider:
            display_fiducia = consumer_data["debtor_name"]
        elif "atas nama bpkb" in provider:
            display_fiducia = consumer_data["name_on_bpkb"]
        else:
            display_fiducia = consumer_data["fiducia_provider"]

        consumer_data["display_fiducia"] = display_fiducia
        consumer_data["firm_name"] = (
            document.notary.firm_name if document and document.notary else ""
        )
        consumer_data["voucher_code"] = document.voucher_code if document else ""
        consumer_data["certificate_number"] = (
            document.certificate_number if document else ""
        )
        consumer_data["certificate_date"] = (
            document.certificate_date if document else ""
        )
        consumer_data["document_id"] = document.id if document else None
        consumer_data["document_number"] = document.document_number if document else ""
        consumer_data["document_date"] = document.document_date if document else None
        consumer_data["consumer_document_created_at"] = consumer_doc_map.get(
            consumer_data["agreement_number"]
        )

    context = {
        "enriched_consumers": consumers_data,
        "tanggal_mulai": tanggal_mulai or "",
        "tanggal_selesai": tanggal_selesai or "",
        "error_message": error_message,
    }
    return render(request, "AktaSertifikat/consumer-list-mu.html", context)


@login_required(login_url="")
@permission_required("update-sertifikat")
def import_excel_sertifikat(request):
    if request.method == "POST":
        try:
            branch_codes = _write_branch_codes_or_forbidden(request.user)
            if branch_codes is None:
                return HttpResponse("Forbidden", status=403)

            file = request.FILES.get("excel_file")
            if not file:
                return HttpResponse("File tidak ditemukan", status=400)

            # Validate Excel file
            validation_error = validate_excel_file_or_error(file)
            if validation_error:
                return validation_error

            wb = load_workbook(filename=file)
            ws = wb.active

            duplicate_vouchers = []

            for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                consumer_id = row[0].value
                if not consumer_id:
                    continue

                try:
                    consumer_qs = Consumer.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        continue
                    consumer.agreement_number = row[1].value or None
                    consumer._skip_signal = True
                    consumer.save()

                    doc = LegalDocuments.objects.filter(consumer=consumer).first()
                    if not doc:
                        continue

                    voucher_code = row[6].value

                    if voucher_code and LegalDocuments.objects.filter(
                        voucher_code=voucher_code
                    ).exclude(id=doc.id).exists():
                        duplicate_vouchers.append(str(voucher_code))
                        continue

                    doc.registration_date = parse_date(row[5].value)
                    doc.voucher_code = voucher_code
                    doc.certificate_number = row[7].value
                    doc.certificate_date = parse_date(row[8].value)
                    doc.certificate_path = row[9].value
                    doc.pnbp_path = row[10].value
                    doc.payment_date = parse_date(row[11].value)
                    doc.save()

                except Consumer.DoesNotExist:
                    continue

            if duplicate_vouchers:
                return HttpResponse(
                    "Import berhasil.\n"
                    "Kode Voucher duplikat tidak terinput:\n"
                    + ", ".join(html.escape(v) for v in duplicate_vouchers),
                    status=200,
                    content_type="text/plain; charset=utf-8"
                )

            return HttpResponse("Import berhasil", status=200, content_type="text/plain; charset=utf-8")

        except Exception as e:
            return HttpResponse(html.escape(str(e)), status=500, content_type="text/plain; charset=utf-8")

    return HttpResponse("Invalid request", status=400)


@login_required(login_url="")
@permission_required("update-sertifikat")
def import_excel_sertifikat_pf(request):
    if request.method == "POST":
        try:
            branch_codes = _write_branch_codes_or_forbidden(request.user)
            if branch_codes is None:
                return HttpResponse("Forbidden", status=403)

            file = request.FILES.get("excel_file")
            if not file:
                return HttpResponse("File tidak ditemukan", status=400)

            # Validate Excel file
            validation_error = validate_excel_file_or_error(file)
            if validation_error:
                return validation_error

            wb = load_workbook(filename=file)
            ws = wb.active

            duplicate_vouchers = []

            for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                consumer_id = row[0].value
                if not consumer_id:
                    continue

                try:
                    consumer_qs = ChangeOfFiduciary.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        continue
                    consumer.agreement_number = row[1].value or None
                    consumer._skip_signal = True
                    consumer.save()

                    doc = LegalDocumentsPF.objects.filter(
                        change_of_fiduciary=consumer
                    ).first()
                    if not doc:
                        continue

                    voucher_code = row[6].value

                    if voucher_code and LegalDocumentsPF.objects.filter(
                        voucher_code=voucher_code
                    ).exclude(id=doc.id).exists():
                        duplicate_vouchers.append(str(voucher_code))
                        continue

                    doc.registration_date = parse_date(row[5].value)
                    doc.voucher_code = voucher_code
                    doc.certificate_number = row[7].value
                    doc.certificate_date = parse_date(row[8].value)
                    doc.certificate_path = row[9].value
                    doc.pnbp_path = row[10].value
                    doc.payment_date = parse_date(row[11].value)
                    doc.save()

                except ChangeOfFiduciary.DoesNotExist:
                    continue

            if duplicate_vouchers:
                return HttpResponse(
                    "Import berhasil.\n"
                    "Kode Voucher duplikat tidak terinput:\n"
                    + ", ".join(html.escape(v) for v in duplicate_vouchers),
                    status=200,
                    content_type="text/plain; charset=utf-8"
                )

            return HttpResponse("Import berhasil", status=200, content_type="text/plain; charset=utf-8")

        except Exception as e:
            return HttpResponse(html.escape(str(e)), status=500, content_type="text/plain; charset=utf-8")

    return HttpResponse("Invalid request", status=400)


@login_required(login_url="")
@permission_required("update-sertifikat")
def import_excel_sertifikat_oc(request):
    if request.method == "POST":
        try:
            branch_codes = _write_branch_codes_or_forbidden(request.user)
            if branch_codes is None:
                return HttpResponse("Forbidden", status=403)

            file = request.FILES.get("excel_file")
            if not file:
                return HttpResponse("File tidak ditemukan", status=400)

            # Validate Excel file
            validation_error = validate_excel_file_or_error(file)
            if validation_error:
                return validation_error

            wb = load_workbook(filename=file)
            ws = wb.active

            duplicate_vouchers = []

            for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                consumer_id = row[0].value
                if not consumer_id:
                    continue

                try:
                    consumer_qs = ConsumerOC.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        continue
                    consumer.agreement_number = row[1].value or None
                    consumer._skip_signal = True
                    consumer.save()

                    doc = LegalDocumentOC.objects.filter(consumer=consumer).first()
                    if not doc:
                        continue

                    voucher_code = row[6].value

                    if voucher_code and LegalDocumentsOC.objects.filter(
                        voucher_code=voucher_code
                    ).exclude(id=doc.id).exists():
                        duplicate_vouchers.append(str(voucher_code))
                        continue

                    doc.registration_date = parse_date(row[5].value)
                    doc.voucher_code = voucher_code
                    doc.certificate_number = row[7].value
                    doc.certificate_date = parse_date(row[8].value)
                    doc.certificate_path = row[9].value
                    doc.pnbp_path = row[10].value
                    doc.payment_date = parse_date(row[11].value)
                    doc.save()

                except ConsumerOC.DoesNotExist:
                    continue

            if duplicate_vouchers:
                return HttpResponse(
                    "Import berhasil.\n"
                    "Kode Voucher duplikat tidak terinput:\n"
                    + ", ".join(html.escape(v) for v in duplicate_vouchers),
                    status=200,
                    content_type="text/plain; charset=utf-8"
                )

            return HttpResponse("Import berhasil", status=200, content_type="text/plain; charset=utf-8")

        except Exception as e:
            return HttpResponse(html.escape(str(e)), status=500, content_type="text/plain; charset=utf-8")

    return HttpResponse("Invalid request", status=400)


@login_required(login_url="")
@permission_required("update-sertifikat")
def import_excel_sertifikat_mu(request):
    if request.method == "POST":
        try:
            branch_codes = _write_branch_codes_or_forbidden(request.user)
            if branch_codes is None:
                return HttpResponse("Forbidden", status=403)

            file = request.FILES.get("excel_file")
            if not file:
                return HttpResponse("File tidak ditemukan", status=400)

            # Validate Excel file
            validation_error = validate_excel_file_or_error(file)
            if validation_error:
                return validation_error

            wb = load_workbook(filename=file)
            ws = wb.active

            duplicate_vouchers = []

            for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                consumer_id = row[0].value
                if not consumer_id:
                    continue

                try:
                    consumer_qs = ConsumerMU.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        continue
                    consumer.agreement_number = row[1].value or None
                    consumer._skip_signal = True
                    consumer.save()

                    doc = LegalDocumentMU.objects.filter(consumer=consumer).first()
                    if not doc:
                        continue

                    voucher_code = row[6].value

                    if voucher_code and LegalDocumentMU.objects.filter(
                        voucher_code=voucher_code
                    ).exclude(id=doc.id).exists():
                        duplicate_vouchers.append(str(voucher_code))
                        continue

                    doc.registration_date = parse_date(row[5].value)
                    doc.voucher_code = voucher_code
                    doc.certificate_number = row[7].value
                    doc.certificate_date = parse_date(row[8].value)
                    doc.certificate_path = row[9].value
                    doc.pnbp_path = row[10].value
                    doc.payment_date = parse_date(row[11].value)
                    doc.save()

                except ConsumerMU.DoesNotExist:
                    continue

            if duplicate_vouchers:
                return HttpResponse(
                    "Import berhasil.\n"
                    "Kode Voucher duplikat tidak terinput:\n"
                    + ", ".join(html.escape(v) for v in duplicate_vouchers),
                    status=200,
                    content_type="text/plain; charset=utf-8"
                )

            return HttpResponse("Import berhasil", status=200, content_type="text/plain; charset=utf-8")

        except Exception as e:
            return HttpResponse(html.escape(str(e)), status=500, content_type="text/plain; charset=utf-8")

    return HttpResponse("Invalid request", status=400)


@login_required(login_url="")
@permission_required("read-finance-reg")
def import_excel_finance(request):
    if request.method == "POST":
        branch_codes = []
        if _user_is_staff_cabang(request.user):
            branch_codes = get_user_branch_codes(request.user)
            if not branch_codes:
                return HttpResponse("Forbidden", status=403)

        file = request.FILES.get("excel_file")
        if not file:
            return JsonResponse({"message": "File tidak ditemukan"}, status=400)

        wb = load_workbook(filename=file, read_only=True)
        ws = wb.active

        updated_count = 0
        skipped_details = []
        failed_details = []

        CERT_COL_INDEX = 8
        PAYMENT_DATE_COL_INDEX = 14

        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                consumer_id = row[0].value
                if not consumer_id:
                    continue

                cert_number_cell_value = row[CERT_COL_INDEX].value
                cert_number_val_upper = str(cert_number_cell_value or '').strip().upper()

                if "CLEAR" in cert_number_val_upper or "GAGAL" in cert_number_val_upper:
                    consumer_qs = Consumer.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        failed_details.append(f"Baris {idx}: Consumer ID '{consumer_id}' tidak ditemukan.")
                        continue

                    doc = LegalDocuments.objects.filter(consumer=consumer).first()
                    if not doc:
                        failed_details.append(f"Baris {idx}: Dokumen legal untuk Consumer ID '{consumer_id}' tidak ditemukan.")
                        continue

                    doc.certificate_number = cert_number_cell_value
                    
                    payment_date_val = parse_date(row[PAYMENT_DATE_COL_INDEX].value)
                    doc.payment_date = payment_date_val

                    doc.save(update_fields=['certificate_number', 'payment_date'])
                    
                    updated_count += 1
                else:
                    reason = f"Status '{cert_number_cell_value}' tidak diizinkan" if cert_number_cell_value else "Status kosong"
                    skipped_details.append(f"Baris {idx}: Dilewati ({reason})")
                    continue

            except Exception as e:
                failed_details.append(f"Baris {idx}: Terjadi error ({str(e)})")
                continue
        
        response_data = {
            "updated_count": updated_count,
            "skipped_count": len(skipped_details),
            "failed_count": len(failed_details),
            "skipped_details": skipped_details[:5],
            "failed_details": failed_details[:5],
        }
        return JsonResponse(response_data)

    return JsonResponse({"message": "Metode request tidak valid"}, status=400)


@login_required(login_url="")
@permission_required("read-update-sertifikat")
def export_excel_sertifikat(request):
    deed_date = request.GET.get("deed_date")
    deed_date_to = request.GET.get("deed_date_to")

    if not deed_date:
        return HttpResponse("Tanggal akta tidak diberikan.", status=400)

    legal_docs = LegalDocuments.objects.all()

    if deed_date and deed_date_to:
        try:
            d1 = datetime.strptime(deed_date, "%Y-%m-%d")
            d2 = datetime.strptime(deed_date_to, "%Y-%m-%d")

            selisih = (d2 - d1).days

            if selisih < 0:
                return HttpResponse("Tanggal tidak valid (dari > sampai).", status=400)

            if selisih > 7:
                messages.error(request, "Range export maksimal 7 hari.")
                return redirect(request.META.get("HTTP_REFERER", "/"))

            legal_docs = legal_docs.filter(document_date__range=[deed_date, deed_date_to])

        except ValueError:
            return HttpResponse("Format tanggal tidak valid.", status=400)

    else:
        legal_docs = legal_docs.filter(document_date=deed_date)

    legal_docs = legal_docs.select_related("consumer", "notary")

    legal_docs = _filter_by_branch(
        legal_docs,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    consumers = Consumer.objects.filter(
        id__in=legal_docs.values_list("consumer_id", flat=True)
    )

    document_map = {doc.consumer_id: doc for doc in legal_docs}

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sertifikat"

    ws.append([
        "ID",
        "No. Perjanjian",
        "No. Akta",
        "Notaris",
        "Nama",
        "Tanggal Daftar",
        "Kode Voucher",
        "No. Sertifikat",
        "Tanggal Sertifikat",
        "Path Sertifikat",
        "Path PNBP",
        "Tanggal Pembayaran",
    ])

    for consumer in consumers:
        doc = document_map.get(consumer.id)

        ws.append([
            consumer.id,
            consumer.agreement_number,
            doc.document_number if doc else "",
            doc.notary.firm_name if doc and doc.notary else "",
            consumer.debtor_name,
            doc.registration_date.strftime("%d-%m-%Y") if doc and doc.registration_date else "",
            doc.voucher_code if doc else "",
            doc.certificate_number if doc else "",
            doc.certificate_date.strftime("%d-%m-%Y") if doc and doc.certificate_date else "",
            doc.certificate_path if doc else "",
            doc.pnbp_path if doc else "",
            doc.payment_date.strftime("%d-%m-%Y") if doc and doc.payment_date else "",
        ])

    if deed_date and deed_date_to:
        filename = f"sertifikat_{deed_date}_to_{deed_date_to}.xlsx"
    else:
        filename = f"sertifikat_{deed_date}.xlsx"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    wb.save(response)
    return response


@login_required(login_url="")
@permission_required("read-update-sertifikat")
def export_excel_sertifikat_pf(request):
    deed_date = request.GET.get("deed_date")
    if not deed_date:
        return HttpResponse("Tanggal akta tidak diberikan.", status=400)

    legal_docs_pf = LegalDocumentsPF.objects.filter(
        document_date=deed_date
    ).select_related("notary", "change_of_fiduciary")
    legal_docs_pf = _filter_by_branch(
        legal_docs_pf,
        request.user,
        [
            "change_of_fiduciary__branch__branch_code",
            "change_of_fiduciary__branch_id",
            "change_of_fiduciary__branch__pk",
        ],
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sertifikat PF"

    ws.append(
        [
            "ID",
            "No. Perjanjian",
            "No. Akta",
            "Notaris",
            "Nama",
            "Tanggal Daftar",
            "Kode Voucher",
            "No. Sertifikat",
            "Tanggal Sertifikat",
            "Path Sertifikat",
            "Path PNBP",
            "Tanggal Pembayaran",
        ]
    )

    for doc in legal_docs_pf:
        consumer = doc.change_of_fiduciary
        notary_name = doc.notary.firm_name if doc.notary else ""
        agreement_number = (
            getattr(consumer, "new_agreement_number", None) or consumer.agreement_number
        )
        debtor_name = getattr(consumer, "new_debtor_name", None) or consumer.debtor_name

        ws.append(
            [
                consumer.id,
                agreement_number,
                doc.document_number,
                notary_name,
                debtor_name,
                (
                    doc.registration_date.strftime("%d-%m-%Y")
                    if doc and doc.registration_date
                    else ""
                ),
                doc.voucher_code if doc else "",
                doc.certificate_number if doc else "",
                (
                    doc.certificate_date.strftime("%d-%m-%Y")
                    if doc and doc.certificate_date
                    else ""
                ),
                doc.certificate_path if doc else "",
                doc.pnbp_path if doc else "",
                (
                    doc.payment_date.strftime("%d-%m-%Y")
                    if doc and doc.payment_date
                    else ""
                ),
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"sertifikat_pf_{deed_date}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response


@login_required(login_url="")
@permission_required("read-update-sertifikat")
def export_excel_sertifikat_oc(request):
    deed_date = request.GET.get("deed_date")
    deed_date_to = request.GET.get("deed_date_to")

    if not deed_date:
        return HttpResponse("Tanggal akta tidak diberikan.", status=400)

    legal_docs = LegalDocumentsOC.objects.all()

    if deed_date and deed_date_to:
        try:
            d1 = datetime.strptime(deed_date, "%Y-%m-%d")
            d2 = datetime.strptime(deed_date_to, "%Y-%m-%d")

            selisih = (d2 - d1).days

            if selisih < 0:
                return HttpResponse("Tanggal tidak valid (dari > sampai).", status=400)

            if selisih > 7:
                messages.error(request, "Range export maksimal 7 hari.")
                return redirect(request.META.get("HTTP_REFERER", "/"))

            legal_docs = legal_docs.filter(document_date__range=[deed_date, deed_date_to])

        except ValueError:
            return HttpResponse("Format tanggal tidak valid.", status=400)

    else:
        legal_docs = legal_docs.filter(document_date=deed_date)

    legal_docs = legal_docs.select_related("consumer", "notary")

    legal_docs = _filter_by_branch(
        legal_docs,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    consumers = ConsumerOC.objects.filter(
        id__in=legal_docs.values_list("consumer_id", flat=True)
    )

    document_map = {doc.consumer_id: doc for doc in legal_docs}

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sertifikat OC"

    ws.append([
        "ID",
        "No. Perjanjian",
        "No. Akta",
        "Notaris",
        "Nama",
        "Tanggal Daftar",
        "Kode Voucher",
        "No. Sertifikat",
        "Tanggal Sertifikat",
        "Path Sertifikat",
        "Path PNBP",
        "Tanggal Pembayaran",
    ])

    for consumer in consumers:
        doc = document_map.get(consumer.id)

        ws.append([
            consumer.id,
            consumer.agreement_number,
            doc.document_number if doc else "",
            doc.notary.firm_name if doc and doc.notary else "",
            consumer.debtor_name,
            doc.registration_date.strftime("%d-%m-%Y") if doc and doc.registration_date else "",
            doc.voucher_code if doc else "",
            doc.certificate_number if doc else "",
            doc.certificate_date.strftime("%d-%m-%Y") if doc and doc.certificate_date else "",
            doc.certificate_path if doc else "",
            doc.pnbp_path if doc else "",
            doc.payment_date.strftime("%d-%m-%Y") if doc and doc.payment_date else "",
        ])

    if deed_date and deed_date_to:
        filename = f"sertifikat_oc_{deed_date}_to_{deed_date_to}.xlsx"
    else:
        filename = f"sertifikat_oc_{deed_date}.xlsx"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    wb.save(response)
    return response


@login_required(login_url="")
@permission_required("read-update-sertifikat")
def export_excel_sertifikat_mu(request):
    deed_date = request.GET.get("deed_date")
    deed_date_to = request.GET.get("deed_date_to")

    if not deed_date:
        return HttpResponse("Tanggal akta tidak diberikan.", status=400)

    legal_docs = LegalDocumentMU.objects.all()

    if deed_date and deed_date_to:
        try:
            d1 = datetime.strptime(deed_date, "%Y-%m-%d")
            d2 = datetime.strptime(deed_date_to, "%Y-%m-%d")

            selisih = (d2 - d1).days

            if selisih < 0:
                return HttpResponse("Tanggal tidak valid (dari > sampai).", status=400)

            if selisih > 7:
                messages.error(request, "Range export maksimal 7 hari.")
                return redirect(request.META.get("HTTP_REFERER", "/"))

            legal_docs = legal_docs.filter(document_date__range=[deed_date, deed_date_to])

        except ValueError:
            return HttpResponse("Format tanggal tidak valid.", status=400)

    else:
        legal_docs = legal_docs.filter(document_date=deed_date)

    legal_docs = legal_docs.select_related("consumer", "notary")

    legal_docs = _filter_by_branch(
        legal_docs,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    consumers = ConsumerMU.objects.filter(
        id__in=legal_docs.values_list("consumer_id", flat=True)
    )

    document_map = {doc.consumer_id: doc for doc in legal_docs}

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sertifikat MU"

    ws.append([
        "ID",
        "No. Perjanjian",
        "No. Akta",
        "Notaris",
        "Nama",
        "Tanggal Daftar",
        "Kode Voucher",
        "No. Sertifikat",
        "Tanggal Sertifikat",
        "Path Sertifikat",
        "Path PNBP",
        "Tanggal Pembayaran",
    ])

    for consumer in consumers:
        doc = document_map.get(consumer.id)

        ws.append([
            consumer.id,
            consumer.agreement_number,
            doc.document_number if doc else "",
            doc.notary.firm_name if doc and doc.notary else "",
            consumer.debtor_name,
            doc.registration_date.strftime("%d-%m-%Y") if doc and doc.registration_date else "",
            doc.voucher_code if doc else "",
            doc.certificate_number if doc else "",
            doc.certificate_date.strftime("%d-%m-%Y") if doc and doc.certificate_date else "",
            doc.certificate_path if doc else "",
            doc.pnbp_path if doc else "",
            doc.payment_date.strftime("%d-%m-%Y") if doc and doc.payment_date else "",
        ])

    if deed_date and deed_date_to:
        filename = f"sertifikat_mu_{deed_date}_to_{deed_date_to}.xlsx"
    else:
        filename = f"sertifikat_mu_{deed_date}.xlsx"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    wb.save(response)
    return response


@login_required(login_url="")
@permission_required("read-finance-reg")
def export_excel_finance(request):
    export_status = (request.GET.get("export_status") or "").lower()
    deed_date = request.GET.get("deed_date")

    if export_status == "check" and not deed_date:
        return HttpResponse("Tanggal akta harus diisi untuk status CHECK.", status=400)

    legal_docs = LegalDocuments.objects.select_related("consumer", "consumer__branch", "notary")

    if deed_date:
        legal_docs = legal_docs.filter(document_date=deed_date)

    legal_docs = _filter_by_branch(
        legal_docs,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    legal_docs = legal_docs.order_by("id")

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Data Sertifikat")

    ws.append(
        [
            "ID Konsumen",
            "Kode Cabang",
            "Nama Cabang",
            "No. Akta",
            "Tanggal Akta",
            "No. Perjanjian",
            "Nama Debitur",
            "Kode Voucher",
            "No. Sertifikat",
            "Tanggal Sertifikat",
            "Nilai Penjaminan",
            "Nominal PNBP",
            "Jasa",
            "Nama Notaris",
            "Tanggal Pembayaran",
            "Status",
        ]
    )

    for doc in legal_docs.iterator(chunk_size=1000):
        doc_item = {
            "voucher_code": doc.voucher_code,
            "registration_date": doc.registration_date,
            "certificate_number": doc.certificate_number,
            "payment_date": doc.payment_date,
            "certificate_path": doc.certificate_path,
            "certificate_date": doc.certificate_date,
        }

        status = determine_status(doc_item).lower()

        if export_status and status != export_status:
            continue

        consumer = doc.consumer

        branch_code = ""
        branch_name = ""

        if consumer:
            if consumer.branch:
                branch_code = consumer.branch.branch_code or ""
                branch_name = consumer.branch.branch_name or ""
            else:
                branch_name = consumer.branch_name or ""

        ws.append(
            [
                consumer.id if consumer else "",
                branch_code,
                branch_name,
                doc.document_number or "",
                doc.document_date,
                consumer.agreement_number if consumer else "",
                consumer.debtor_name if consumer else "",
                doc.voucher_code or "",
                doc.certificate_number or "",
                doc.certificate_date,
                consumer.guarantee_value if consumer else "",
                consumer.pnbp_fee if consumer else "",
                consumer.service_fee if consumer else "",
                doc.notary.firm_name if doc.notary else "",
                doc.payment_date,
                status.upper(),
            ]
        )

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    if export_status == "check" and deed_date:
        filename = f"sertifikat_{deed_date}_status_check.xlsx"
    else:
        filename = f"sertifikat_status_{export_status or 'all'}.xlsx"

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required(login_url="")
@permission_required("read-log-akta")
def akta_minuta_log_list(request):
    # query = request.GET.get("q", "")
    logs = AktaMinutaLog.objects.all()

    if _user_is_staff_cabang(request.user):
        branch_codes = get_user_branch_codes(request.user)
        if not branch_codes:
            logs = logs.none()
        else:
            allowed_agreements = _agreement_numbers_for_branches(Consumer, branch_codes)
            logs = logs.filter(agreement_number__in=allowed_agreements)

    # if query:
    #     logs = logs.filter(agreement_number__icontains=query)

    # logs = logs.order_by("-id")

    # # paginator = Paginator(logs, 10)
    # # page_number = request.GET.get('page')
    # # page_obj = paginator.get_page(page_number)

    # context = {
    #     # 'page_obj': page_obj,
    #     "logs": logs,
    #     "query": query,
    # }
    return render(request, "AktaSertifikat/akta-minuta-log.html")

@login_required(login_url="")
@permission_required("read-log-akta")
def akta_minuta_log_api(request):
    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 10))
    search_value = request.GET.get('search[value]', '').strip()

    queryset = AktaMinutaLog.objects.select_related('pic').order_by('-print_date')

    if search_value:
        queryset = queryset.filter(
            Q(agreement_number__icontains=search_value) |
            Q(pic__name__icontains=search_value)
        )

    records_total = AktaMinutaLog.objects.count()
    records_filtered = queryset.count()

    logs = queryset[start:start + length]

    data = []
    for i, log in enumerate(logs, start=1):
        data.append([
            start + i,
            log.agreement_number,
            log.pic.name if log.pic else '-',
            log.print_date.strftime('%Y-%m-%d %H:%M:%S') if log.print_date else '-'
        ])

    return JsonResponse({
        'draw': draw,
        'recordsTotal': records_total,
        'recordsFiltered': records_filtered,
        'data': data
    })

@login_required(login_url="")
@permission_required("read-log-akta-pf")
def akta_minuta_log_list_pf(request):
    # query = request.GET.get("q", "")
    logs = AktaMinutaLogPF.objects.all()

    if _user_is_staff_cabang(request.user):
        branch_codes = get_user_branch_codes(request.user)
        if not branch_codes:
            logs = logs.none()
        else:
            allowed_agreements = _agreement_numbers_for_branches(ChangeOfFiduciary, branch_codes)
            logs = logs.filter(agreement_number__in=allowed_agreements)

    # if query:
    #     logs = logs.filter(agreement_number__icontains=query)

    # logs = logs.order_by("-id")

    # context = {
    #     "logs": logs,
    #     "query": query,
    # }
    return render(request, "AktaSertifikat/akta-minuta-log-pf.html")

@login_required(login_url="")
@permission_required("read-log-akta-pf")
def akta_minuta_log_api_pf(request):
    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 10))
    search_value = request.GET.get('search[value]', '').strip()

    queryset = AktaMinutaLogPF.objects.select_related('pic').order_by('-print_date')

    if search_value:
        queryset = queryset.filter(
            Q(agreement_number__icontains=search_value) |
            Q(pic__name__icontains=search_value)
        )

    records_total = AktaMinutaLogPF.objects.count()
    records_filtered = queryset.count()

    logs = queryset[start:start + length]

    data = []
    for i, log in enumerate(logs, start=1):
        data.append([
            start + i,
            log.agreement_number,
            log.pic.name if log.pic else '-',
            log.print_date.strftime('%Y-%m-%d %H:%M:%S') if log.print_date else '-'
        ])

    return JsonResponse({
        'draw': draw,
        'recordsTotal': records_total,
        'recordsFiltered': records_filtered,
        'data': data
    })

@login_required(login_url="")
@permission_required("read-log-akta-oc")
def akta_minuta_log_list_oc(request):
    # query = request.GET.get("q", "")
    logs = AktaMinutaLogOC.objects.all()

    if _user_is_staff_cabang(request.user):
        branch_codes = get_user_branch_codes(request.user)
        if not branch_codes:
            logs = logs.none()
        else:
            allowed_agreements = _agreement_numbers_for_branches(ConsumerOC, branch_codes)
            logs = logs.filter(agreement_number__in=allowed_agreements)

    # if query:
    #     logs = logs.filter(agreement_number__icontains=query)

    # logs = logs.order_by("-id")

    # context = {
    #     "logs": logs,
    #     "query": query,
    # }
    return render(request, "AktaSertifikat/akta-minuta-log-oc.html")

@login_required(login_url="")
@permission_required("read-log-akta-oc")
def akta_minuta_log_api_oc(request):
    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 10))
    search_value = request.GET.get('search[value]', '').strip()

    queryset = AktaMinutaLogOC.objects.select_related('pic').order_by('-print_date')

    if search_value:
        queryset = queryset.filter(
            Q(agreement_number__icontains=search_value) |
            Q(pic__name__icontains=search_value)
        )

    records_total = AktaMinutaLogOC.objects.count()
    records_filtered = queryset.count()

    logs = queryset[start:start + length]

    data = []
    for i, log in enumerate(logs, start=1):
        data.append([
            start + i,
            log.agreement_number,
            log.pic.name if log.pic else '-',
            log.print_date.strftime('%Y-%m-%d %H:%M:%S') if log.print_date else '-'
        ])

    return JsonResponse({
        'draw': draw,
        'recordsTotal': records_total,
        'recordsFiltered': records_filtered,
        'data': data
    })

@login_required(login_url="")
@permission_required("read-log-akta-mu")
def akta_minuta_log_list_mu(request):
    # query = request.GET.get("q", "")
    logs = AktaMinutaLogMU.objects.all()

    if _user_is_staff_cabang(request.user):
        branch_codes = get_user_branch_codes(request.user)
        if not branch_codes:
            logs = logs.none()
        else:
            allowed_agreements = _agreement_numbers_for_branches(ConsumerMU, branch_codes)
            logs = logs.filter(agreement_number__in=allowed_agreements)

    # if query:
    #     logs = logs.filter(agreement_number__icontains=query)

    # logs = logs.order_by("-id")

    # context = {
    #     "logs": logs,
    #     "query": query,
    # }
    return render(request, "AktaSertifikat/akta-minuta-log-mu.html")

@login_required(login_url="")
@permission_required("read-log-akta-mu")
def akta_minuta_log_api_mu(request):
    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 10))
    search_value = request.GET.get('search[value]', '').strip()

    queryset = AktaMinutaLogMU.objects.select_related('pic').order_by('-print_date')

    if search_value:
        queryset = queryset.filter(
            Q(agreement_number__icontains=search_value) |
            Q(pic__name__icontains=search_value)
        )

    records_total = AktaMinutaLogMU.objects.count()
    records_filtered = queryset.count()

    logs = queryset[start:start + length]

    data = []
    for i, log in enumerate(logs, start=1):
        data.append([
            start + i,
            log.agreement_number,
            log.pic.name if log.pic else '-',
            log.print_date.strftime('%Y-%m-%d %H:%M:%S') if log.print_date else '-'
        ])

    return JsonResponse({
        'draw': draw,
        'recordsTotal': records_total,
        'recordsFiltered': records_filtered,
        'data': data
    })

@login_required(login_url="")
@permission_required("update-akser")
@csrf_protect
def upload_akta_minuta_log(request):
    if request.method == "POST":
        file = request.FILES.get("excel_file")
        if not file:
            return HttpResponse("File tidak ditemukan", status=400)

        allowed_agreements = None
        if _user_is_staff_cabang(request.user):
            branch_codes = get_user_branch_codes(request.user)
            if not branch_codes:
                return HttpResponse("Forbidden", status=403)
            allowed_agreements = _agreement_numbers_for_branches(Consumer, branch_codes)

        try:
            wb = load_workbook(filename=file)
            ws = wb.active
        except Exception as e:
            return HttpResponse(f"File Excel tidak valid: {html.escape(str(e))}", status=400, content_type="text/plain; charset=utf-8")

        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            agreement_number = row[0].value

            if not agreement_number:
                print(f"[Baris {idx}] Nomor Perjanjian kosong. Lewati.")
                continue

            if allowed_agreements is not None and agreement_number not in allowed_agreements:
                continue

            if AktaMinutaLog.objects.filter(agreement_number=agreement_number).exists():
                print(f"[Baris {idx}] Nomor Perjanjian {agreement_number} sudah ada. Lewati.")
                continue

            AktaMinutaLog.objects.create(
                agreement_number=agreement_number,
                pic=request.user,
                print_date=timezone.now(),
            )

        return redirect("akta-minuta-log-list")

    return HttpResponse("Invalid request method", status=400)

@login_required(login_url="")
@permission_required("update-akser-oc")
@csrf_protect
def upload_akta_minuta_log_oc(request):
    if request.method == "POST":
        file = request.FILES.get("excel_file")
        if not file:
            return HttpResponse("File tidak ditemukan", status=400)

        allowed_agreements = None
        if _user_is_staff_cabang(request.user):
            branch_codes = get_user_branch_codes(request.user)
            if not branch_codes:
                return HttpResponse("Forbidden", status=403)
            allowed_agreements = _agreement_numbers_for_branches(ConsumerOC, branch_codes)

        try:
            wb = load_workbook(filename=file)
            ws = wb.active
        except Exception as e:
            return HttpResponse(f"File Excel tidak valid: {html.escape(str(e))}", status=400, content_type="text/plain; charset=utf-8")

        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            agreement_number = row[0].value

            if not agreement_number:
                print(f"[Baris {idx}] Nomor Perjanjian kosong. Lewati.")
                continue

            if allowed_agreements is not None and agreement_number not in allowed_agreements:
                continue
            
            if AktaMinutaLogOC.objects.filter(agreement_number=agreement_number).exists():
                print(f"[Baris {idx}] Nomor Perjanjian {agreement_number} sudah ada. Lewati.")
                continue

            AktaMinutaLogOC.objects.create(
                agreement_number=agreement_number,
                pic=request.user,
                print_date=timezone.now(),
            )

        return redirect("akta-minuta-log-list-oc")

    return HttpResponse("Invalid request method", status=400)

@login_required(login_url="")
@permission_required("update-akser-mu")
@csrf_protect
def upload_akta_minuta_log_mu(request):
    if request.method == "POST":
        file = request.FILES.get("excel_file")
        if not file:
            return HttpResponse("File tidak ditemukan", status=400)

        allowed_agreements = None
        if _user_is_staff_cabang(request.user):
            branch_codes = get_user_branch_codes(request.user)
            if not branch_codes:
                return HttpResponse("Forbidden", status=403)
            allowed_agreements = _agreement_numbers_for_branches(ConsumerMU, branch_codes)

        try:
            wb = load_workbook(filename=file)
            ws = wb.active
        except Exception as e:
            return HttpResponse(f"File Excel tidak valid: {html.escape(str(e))}", status=400, content_type="text/plain; charset=utf-8")

        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            agreement_number = row[0].value

            if not agreement_number:
                print(f"[Baris {idx}] Nomor Perjanjian kosong. Lewati.")
                continue

            if allowed_agreements is not None and agreement_number not in allowed_agreements:
                continue
            
            if AktaMinutaLogMU.objects.filter(agreement_number=agreement_number).exists():
                print(f"[Baris {idx}] Nomor Perjanjian {agreement_number} sudah ada. Lewati.")
                continue

            AktaMinutaLogMU.objects.create(
                agreement_number=agreement_number,
                pic=request.user,
                print_date=timezone.now(),
            )

        return redirect("akta-minuta-log-list-mu")

    return HttpResponse("Invalid request method", status=400)

@login_required(login_url="")
@permission_required("update-akser-pf")
@csrf_protect
def upload_akta_minuta_log_pf(request):
    if request.method == "POST":
        file = request.FILES.get("excel_file")
        if not file:
            return HttpResponse("File tidak ditemukan", status=400)

        allowed_agreements = None
        if _user_is_staff_cabang(request.user):
            branch_codes = get_user_branch_codes(request.user)
            if not branch_codes:
                return HttpResponse("Forbidden", status=403)
            allowed_agreements = _agreement_numbers_for_branches(ChangeOfFiduciary, branch_codes)

        try:
            wb = load_workbook(filename=file)
            ws = wb.active
        except Exception as e:
            return HttpResponse(f"File Excel tidak valid: {html.escape(str(e))}", status=400, content_type="text/plain; charset=utf-8")

        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            agreement_number = row[0].value

            if not agreement_number:
                print(f"[Baris {idx}] Nomor Perjanjian kosong. Lewati.")
                continue

            if allowed_agreements is not None and agreement_number not in allowed_agreements:
                continue
            
            if AktaMinutaLogPF.objects.filter(agreement_number=agreement_number).exists():
                print(f"[Baris {idx}] Nomor Perjanjian {agreement_number} sudah ada. Lewati.")
                continue

            AktaMinutaLogPF.objects.create(
                agreement_number=agreement_number,
                pic=request.user,
                print_date=timezone.now(),
            )

        return redirect("akta-minuta-log-list-pf")

    return HttpResponse("Invalid request method", status=400)

@login_required(login_url="")
@permission_required("read-akser")
def export_excel_akser(request):
    tanggal_mulai = request.GET.get("tanggal_mulai")
    tanggal_selesai = request.GET.get("tanggal_selesai")

    if not tanggal_mulai or not tanggal_selesai:
        return HttpResponse(
            "Tanggal mulai dan tanggal selesai wajib diisi.", status=400
        )

    start_date = parse_date(tanggal_mulai)
    end_date = parse_date(tanggal_selesai)

    if not start_date or not end_date:
        return HttpResponse("Format tanggal tidak valid.", status=400)

    if (end_date - start_date).days > 31:
        return HttpResponse("Rentang tanggal maksimal 31 hari.", status=400)

    from import_document.models import ConsumerDocument
    legal_docs = LegalDocuments.objects.select_related(
        "consumer", "consumer__branch", "notary"
    ).filter(
        document_date__range=(start_date, end_date)
    ).distinct()
    legal_docs = _filter_by_branch(
        legal_docs,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    legal_docs = legal_docs.order_by("document_date", "certificate_date", "id")
    agreement_list = list(legal_docs.values_list('consumer__agreement_number', flat=True))
    cd_created_map = dict(
        ConsumerDocument.objects.filter(agreement_number__in=agreement_list)
        .values_list('agreement_number', 'created_at')
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Akta & Sertifikat"

    ws.append(
        [
            "No. Perjanjian",
            "Nama Cabang",
            "Notaris",
            "Nama Debitur",
            "Tanggal Akta",
            "No Akta",
            "Kode Voucher",
            "No. Sertifikat",
            "Tanggal Sertifikat",
            "Tgl Upload Excel",
            "Tgl Upload Dokumen",
        ]
    )

    for doc in legal_docs:
        consumer = getattr(doc, "consumer", None)
        branch = getattr(consumer, "branch", None)
        notary = getattr(doc, "notary", None)
        agreement_number = getattr(consumer, "agreement_number", "") or ""

        tgl_upload_excel = (
            consumer.created_at.strftime("%d/%m/%Y")
            if getattr(consumer, "created_at", None)
            else ""
        )

        cd_created_at = cd_created_map.get(agreement_number)
        tgl_upload = cd_created_at.strftime("%d/%m/%Y") if cd_created_at else ""

        ws.append(
            [
                agreement_number,
                getattr(branch, "branch_name", "") or "",
                getattr(notary, "firm_name", "") or "",
                getattr(consumer, "debtor_name", "") or "",
                (
                    doc.document_date.strftime("%d/%m/%Y")
                    if getattr(doc, "document_date", None)
                    else ""
                ),
                getattr(doc, "document_number", "") or "",
                getattr(doc, "voucher_code", "") or "",
                getattr(doc, "certificate_number", "") or "",
                (
                    doc.certificate_date.strftime("%d/%m/%Y")
                    if getattr(doc, "certificate_date", None)
                    else ""
                ),
                tgl_upload_excel,
                tgl_upload,
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = (
        f"akser_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
    )
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response


@login_required(login_url="")
def export_excel_akser_pf(request):
    if not _can_export_akser_pf(request.user):
        return render(request, "403.html", status=403)

    tanggal_mulai = request.GET.get("tanggal_mulai")
    tanggal_selesai = request.GET.get("tanggal_selesai")

    if not tanggal_mulai or not tanggal_selesai:
        return HttpResponse("Tanggal mulai dan tanggal selesai wajib diisi.", status=400)

    start_date = parse_date(tanggal_mulai)
    end_date = parse_date(tanggal_selesai)

    if not start_date or not end_date:
        return HttpResponse("Format tanggal tidak valid.", status=400)

    if (end_date - start_date).days > 31:
        return HttpResponse("Rentang tanggal maksimal 31 hari.", status=400)

    docs_q = Q(change_of_fiduciary__created_at__date__range=(start_date, end_date))

    legal_docs_pf = LegalDocumentsPF.objects.select_related(
        "change_of_fiduciary", "change_of_fiduciary__branch", "notary"
    ).filter(docs_q)
    legal_docs_pf = _filter_by_branch(
        legal_docs_pf,
        request.user,
        [
            "change_of_fiduciary__branch__branch_code",
            "change_of_fiduciary__branch_id",
            "change_of_fiduciary__branch__pk",
        ],
    )

    legal_docs_pf = legal_docs_pf.order_by("document_date", "certificate_date", "id")

    from import_document_pf.models import ConsumerDocumentPF
    agreement_list = list(
        legal_docs_pf.values_list(
            "change_of_fiduciary__new_agreement_number",
            "change_of_fiduciary__agreement_number",
        )
    )
    agree_set = set()
    for new_ag, old_ag in agreement_list:
        if new_ag:
            agree_set.add(new_ag)
        if old_ag:
            agree_set.add(old_ag)

    cdpf_created_map = {}
    if agree_set:
        cdpf_created_map = dict(
            ConsumerDocumentPF.objects.filter(agreement_number__in=agree_set)
            .values_list("agreement_number", "created_at")
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sertifikat PF"

    ws.append(
        [
            "No. Perjanjian",
            "Nama Cabang",
            "Notaris",
            "Nama Debitur",
            "Tanggal Akta",
            "No Akta",
            "Kode Voucher",
            "No. Sertifikat",
            "Tanggal Sertifikat",
            "Tgl Upload Excel",
            "Tgl Upload Dokumen",
        ]
    )

    for doc in legal_docs_pf:
        cof = getattr(doc, "change_of_fiduciary", None)
        branch = getattr(cof, "branch", None)
        notary = getattr(doc, "notary", None)

        agreement_number = ""
        debtor_name = ""
        if cof:
            agreement_number = (
                getattr(cof, "new_agreement_number", None)
                or getattr(cof, "agreement_number", "")
                or ""
            )
            debtor_name = (
                getattr(cof, "new_debtor_name", None)
                or getattr(cof, "debtor_name", "")
                or ""
            )

        tgl_upload_excel = (
            cof.created_at.strftime("%d/%m/%Y") if getattr(cof, "created_at", None) else ""
        )

        cdpf_created_at = cdpf_created_map.get(agreement_number)
        tgl_upload_dokumen = cdpf_created_at.strftime("%d/%m/%Y") if cdpf_created_at else ""

        ws.append(
            [
                agreement_number,
                getattr(branch, "branch_name", "") or "",
                getattr(notary, "firm_name", "") or "",
                debtor_name,
                (
                    doc.document_date.strftime("%d/%m/%Y")
                    if getattr(doc, "document_date", None)
                    else ""
                ),
                getattr(doc, "document_number", "") or "",
                getattr(doc, "voucher_code", "") or "",
                getattr(doc, "certificate_number", "") or "",
                (
                    doc.certificate_date.strftime("%d/%m/%Y")
                    if getattr(doc, "certificate_date", None)
                    else ""
                ),
                tgl_upload_excel,
                tgl_upload_dokumen,
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"akser_pf_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response


@login_required(login_url="")
@permission_required("read-akser-mu")
def export_excel_akser_mu(request):
    tanggal_mulai = request.GET.get("tanggal_mulai")
    tanggal_selesai = request.GET.get("tanggal_selesai")

    if not tanggal_mulai or not tanggal_selesai:
        return HttpResponse(
            "Tanggal mulai dan tanggal selesai wajib diisi.", status=400
        )

    start_date = parse_date(tanggal_mulai)
    end_date = parse_date(tanggal_selesai)

    if not start_date or not end_date:
        return HttpResponse("Format tanggal tidak valid.", status=400)

    if start_date > end_date:
        return HttpResponse(
            "Tanggal mulai tidak boleh lebih besar dari tanggal selesai.", status=400
        )

    if (end_date - start_date).days > 31:
        return HttpResponse("Rentang tanggal maksimal 31 hari.", status=400)

    tz = timezone.get_current_timezone()
    start_dt_local = datetime.combine(start_date, time.min)
    end_dt_local = datetime.combine(end_date + timedelta(days=1), time.min)

    start_dt = timezone.make_aware(start_dt_local, tz)
    end_dt = timezone.make_aware(end_dt_local, tz)
    
    legal_docs_mu = LegalDocumentMU.objects.select_related(
        "consumer", "consumer__branch", "notary"
    ).filter(
        consumer__created_at__date__range=(start_date, end_date)
    ).distinct()
    legal_docs_mu = _filter_by_branch(
        legal_docs_mu,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    legal_docs_mu = legal_docs_mu.order_by(
        "consumer__agreement_number", "document_date", "certificate_date", "id"
    )
    
    agreement_list = list(legal_docs_mu.values_list('consumer__agreement_number', flat=True))
    cd_created_map = dict(
        ConsumerDocumentMU.objects.filter(agreement_number__in=agreement_list)
        .values_list('agreement_number', 'created_at')
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Akta & Sertifikat MU"

    ws.append(
        [
            "No. Perjanjian",
            "Nama Cabang",
            "Notaris",
            "Nama Debitur",
            "Tanggal Akta",
            "No Akta",
            "Kode Voucher",
            "No. Sertifikat",
            "Tanggal Sertifikat",
            "Tgl Upload Excel",
            "Tgl Upload Dokumen",
        ]
    )

    for doc in legal_docs_mu:
        consumer = getattr(doc, "consumer", None)
        branch = getattr(consumer, "branch", None)
        notary = getattr(doc, "notary", None)
        agreement_number = getattr(consumer, "agreement_number", "") or ""

        tgl_upload_excel = (
            consumer.created_at.strftime("%d/%m/%Y")
            if getattr(consumer, "created_at", None)
            else ""
        )

        cd_created_at = cd_created_map.get(agreement_number)
        tgl_upload = cd_created_at.strftime("%d/%m/%Y") if cd_created_at else ""

        ws.append(
            [
                agreement_number,
                getattr(branch, "branch_name", "") or "",
                getattr(notary, "firm_name", "") or "",
                getattr(consumer, "debtor_name", "") or "",
                (
                    doc.document_date.strftime("%d/%m/%Y")
                    if getattr(doc, "document_date", None)
                    else ""
                ),
                getattr(doc, "document_number", "") or "",
                getattr(doc, "voucher_code", "") or "",
                getattr(doc, "certificate_number", "") or "",
                (
                    doc.certificate_date.strftime("%d/%m/%Y")
                    if getattr(doc, "certificate_date", None)
                    else ""
                ),
                tgl_upload_excel,
                tgl_upload,
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = (
        f"akser_mu_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
    )
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response


@login_required(login_url="")
@permission_required("read-akser")
def export_excel_akser_oc(request):
    tanggal_mulai = request.GET.get("tanggal_mulai")
    tanggal_selesai = request.GET.get("tanggal_selesai")

    if not tanggal_mulai or not tanggal_selesai:
        return HttpResponse(
            "Tanggal mulai dan tanggal selesai wajib diisi.", status=400
        )

    start_date = parse_date(tanggal_mulai)
    end_date = parse_date(tanggal_selesai)

    if not start_date or not end_date:
        return HttpResponse("Format tanggal tidak valid.", status=400)

    if (end_date - start_date).days > 31:
        return HttpResponse("Rentang tanggal maksimal 31 hari.", status=400)
    
    from import_document_oc.models import ConsumerDocumentOC
    legal_docs_oc = LegalDocumentsOC.objects.select_related(
        "consumer", "consumer__branch", "notary"
    ).filter(
        consumer__created_at__date__range=(start_date, end_date)
    ).distinct()
    legal_docs_oc = _filter_by_branch(
        legal_docs_oc,
        request.user,
        ["consumer__branch__branch_code", "consumer__branch_id", "consumer__branch__pk"],
    )

    agreement_list = list(legal_docs_oc.values_list('consumer__agreement_number', flat=True))
    cd_created_map = dict(
        ConsumerDocumentOC.objects.filter(agreement_number__in=agreement_list)
        .values_list('agreement_number', 'created_at')
    )

    # Urutkan
    legal_docs_oc = legal_docs_oc.order_by(
        "consumer__agreement_number", "document_date", "certificate_date", "id"
    )

    # Buat workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Akta & Sertifikat OC"

    # Header
    ws.append(
        [
            "No. Perjanjian",
            "Nama Cabang",
            "Notaris",
            "Nama Debitur",
            "Tanggal Akta",
            "No Akta",
            "Kode Voucher",
            "No. Sertifikat",
            "Tanggal Sertifikat",
            "Tgl Upload Excel",
            "Tgl Upload Dokumen",
        ]
    )

    # Isi data
    for doc in legal_docs_oc:
        consumer = getattr(doc, "consumer", None)
        branch = getattr(consumer, "branch", None)
        notary = getattr(doc, "notary", None)
        agreement_number = getattr(consumer, "agreement_number", "") or ""

        tgl_upload_excel = (
            consumer.created_at.strftime("%d/%m/%Y")
            if getattr(consumer, "created_at", None)
            else ""
        )

        cd_created_at = cd_created_map.get(agreement_number)
        tgl_upload = cd_created_at.strftime("%d/%m/%Y") if cd_created_at else ""

        ws.append(
            [
                agreement_number,
                getattr(branch, "branch_name", "") or "",
                getattr(notary, "firm_name", "") or "",
                getattr(consumer, "debtor_name", "") or "",
                (
                    doc.document_date.strftime("%d/%m/%Y")
                    if getattr(doc, "document_date", None)
                    else ""
                ),
                getattr(doc, "document_number", "") or "",
                getattr(doc, "voucher_code", "") or "",
                getattr(doc, "certificate_number", "") or "",
                (
                    doc.certificate_date.strftime("%d/%m/%Y")
                    if getattr(doc, "certificate_date", None)
                    else ""
                ),
                tgl_upload_excel,
                tgl_upload,
            ]
        )

    # Response
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = (
        f"akser_oc_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
    )
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response


@login_required(login_url="")
@permission_required("update-sertifikat")
def update_sertifikat_bulk(request):
    if request.method == "POST":
        try:
            branch_codes = _write_branch_codes_or_forbidden(request.user)
            if branch_codes is None:
                return JsonResponse(
                    {"status": "error", "message": "Forbidden"}, status=403
                )

            updates = json.loads(request.body.decode("utf-8"))
            for row in updates:
                consumer_id = row.get("id")
                fields = row.get("fields", {})

                try:
                    consumer_qs = Consumer.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        continue
                    doc = LegalDocuments.objects.filter(consumer=consumer).first()
                    if not doc:
                        continue

                    for field, value in fields.items():
                        if field in [
                            "registration_date",
                            "certificate_date",
                            "payment_date",
                        ]:
                            setattr(doc, field, parse_date(value) if value else None)
                        elif field in [
                            "voucher_code",
                            "certificate_number",
                            "certificate_path",
                            "pnbp_path",
                        ]:
                            setattr(doc, field, value)

                    doc.save()

                except Consumer.DoesNotExist:
                    continue

            return JsonResponse({"status": "success"})

        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)}, status=500)

    return JsonResponse({"status": "error", "message": "Invalid request"}, status=400)


@login_required(login_url="")
@permission_required("update-sertifikat")
def update_sertifikat_bulk_pf(request):
    if request.method == "POST":
        try:
            branch_codes = _write_branch_codes_or_forbidden(request.user)
            if branch_codes is None:
                return JsonResponse(
                    {"status": "error", "message": "Forbidden"}, status=403
                )

            updates = json.loads(request.body.decode("utf-8"))
            for row in updates:
                consumer_id = row.get("id")
                fields = row.get("fields", {})

                try:
                    consumer_qs = ChangeOfFiduciary.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        continue
                    doc = LegalDocumentsPF.objects.filter(
                        change_of_fiduciary=consumer
                    ).first()
                    if not doc:
                        continue

                    for field, value in fields.items():
                        if field in [
                            "registration_date",
                            "certificate_date",
                            "payment_date",
                        ]:
                            setattr(doc, field, parse_date(value) if value else None)
                        elif field in [
                            "voucher_code",
                            "certificate_number",
                            "certificate_path",
                            "pnbp_path",
                        ]:
                            setattr(doc, field, value)

                    doc.save()

                except ChangeOfFiduciary.DoesNotExist:
                    continue

            return JsonResponse({"status": "success"})

        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)}, status=500)

    return JsonResponse({"status": "error", "message": "Invalid request"}, status=400)


@login_required(login_url="")
@permission_required("update-sertifikat")
def update_sertifikat_bulk_oc(request):
    if request.method == "POST":
        try:
            branch_codes = _write_branch_codes_or_forbidden(request.user)
            if branch_codes is None:
                return JsonResponse(
                    {"status": "error", "message": "Forbidden"}, status=403
                )

            updates = json.loads(request.body.decode("utf-8"))
            for row in updates:
                consumer_id = row.get("id")
                fields = row.get("fields", {})

                try:
                    consumer_qs = ConsumerOC.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        continue
                    doc = LegalDocumentOC.objects.filter(consumer=consumer).first()
                    if not doc:
                        continue

                    for field, value in fields.items():
                        if field in [
                            "registration_date",
                            "certificate_date",
                            "payment_date",
                        ]:
                            setattr(doc, field, parse_date(value) if value else None)
                        elif field in [
                            "voucher_code",
                            "certificate_number",
                            "certificate_path",
                            "pnbp_path",
                        ]:
                            setattr(doc, field, value)

                    doc.save()

                except ConsumerOC.DoesNotExist:
                    continue

            return JsonResponse({"status": "success"})

        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)}, status=500)

    return JsonResponse({"status": "error", "message": "Invalid request"}, status=400)


@login_required(login_url="")
@permission_required("update-sertifikat")
def update_sertifikat_bulk_mu(request):
    if request.method == "POST":
        try:
            branch_codes = _write_branch_codes_or_forbidden(request.user)
            if branch_codes is None:
                return JsonResponse(
                    {"status": "error", "message": "Forbidden"}, status=403
                )

            updates = json.loads(request.body.decode("utf-8"))
            for row in updates:
                consumer_id = row.get("id")
                fields = row.get("fields", {})

                try:
                    consumer_qs = ConsumerMU.objects.filter(id=consumer_id)
                    if branch_codes:
                        consumer_qs = consumer_qs.filter(
                            Q(branch__branch_code__in=branch_codes) | Q(branch_id__in=branch_codes)
                        )
                    consumer = consumer_qs.first()
                    if not consumer:
                        continue
                    doc = LegalDocumentMU.objects.filter(consumer=consumer).first()
                    if not doc:
                        continue

                    for field, value in fields.items():
                        if field in [
                            "registration_date",
                            "certificate_date",
                            "payment_date",
                        ]:
                            setattr(doc, field, parse_date(value) if value else None)
                        elif field in [
                            "voucher_code",
                            "certificate_number",
                            "certificate_path",
                            "pnbp_path",
                        ]:
                            setattr(doc, field, value)

                    doc.save()

                except ConsumerMU.DoesNotExist:
                    continue

            return JsonResponse({"status": "success"})

        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)}, status=500)

    return JsonResponse({"status": "error", "message": "Invalid request"}, status=400)


# Helper
def parse_date(value):
    if isinstance(value, datetime):
        return value.date()

    # Jika serial Excel number
    if isinstance(value, (int, float)):
        try:
            dt = pd.to_datetime(value, unit="d", origin="1899-12-30")
            return dt.date()
        except Exception:
            pass

    # Jika string, coba beberapa format umum
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except Exception:
                continue

    return None


def parse_time(value):
    if isinstance(value, time):
        return value

    if isinstance(value, datetime):
        return value.time()

    if isinstance(value, (int, float)):
        try:
            dt = pd.to_datetime(value, unit="d", origin="1899-12-30")
            return dt.time()
        except Exception:
            pass

    if isinstance(value, str):
        for fmt in (
            "%H:%M",
            "%H:%M:%S",
            "%I:%M %p",
            "%I:%M:%S %p",
        ):  # Format 24-jam dan 12-jam
            try:
                return datetime.strptime(value.strip(), fmt).time()
            except Exception:
                continue

    return None


def determine_status(doc_item):
    if not doc_item:
        return "Draft"

    def _has(text, needle):
        return bool(text) and needle in str(text).upper()

    def _is_valid_cert(num):
        if not num:
            return False
        up = str(num).upper()
        return not any(x in up for x in ["CLEAR", "RECHECK", "PENDING", "GAGAL"])

    has_check_info = bool(doc_item["voucher_code"]) and bool(
        doc_item["registration_date"]
    )
    has_recheck_info = has_check_info and _has(
        doc_item["certificate_number"], "RECHECK"
    )
    has_clear_info = has_check_info and _has(
        doc_item["certificate_number"], "CLEAR"
    )
    has_pending_info = has_check_info and _has(
        doc_item["certificate_number"], "PENDING"
    )
    has_gagal_info = has_check_info and _has(
        doc_item["certificate_number"], "GAGAL"
    )
    has_paid_info = has_check_info and bool(doc_item["payment_date"])

    has_done_info = (
        has_paid_info
        and bool(doc_item["certificate_path"])
        and bool(doc_item["certificate_date"])
        and _is_valid_cert(doc_item["certificate_number"])
    )

    if has_done_info:
        return "Done"
    elif has_paid_info:
        return "Paid"
    elif has_gagal_info:
        return "Gagal"
    elif has_clear_info:
        return "Clear"
    elif has_recheck_info:
        return "Recheck"
    elif has_pending_info:
        return "Pending"
    elif has_check_info:
        return "Check"
    return "Draft"


def determine_status_pf(consumer, doc):
    has_on_process_info = (
        doc
        and doc.voucher_code
        and doc.registration_date
        and doc.certificate_number
        and doc.certificate_date
    )

    has_queue_info = (
        has_on_process_info
        and doc.document_number
        and doc.document_date
        and doc.document_time
        and consumer.deed_date
        and consumer.deed_number
    )

    has_done_info = has_queue_info and doc.certificate_path

    if has_done_info:
        return "Done"
    elif has_queue_info:
        return "Queue"
    elif has_on_process_info:
        return "On Process"
    return "Draft"

def extract_no_perjanjian(filename):
    if not filename:
        return None
    return os.path.basename(filename).split('_')[0].strip()

def check_missing_files(doc_model, field_name, folder_path):
    existing_files = set(os.listdir(folder_path))

    docs = doc_model.objects.exclude(**{f"{field_name}__isnull": True}).exclude(**{field_name: ""})

    missing = []

    for doc in docs:
        file_path = getattr(doc, field_name)
        filename = os.path.basename(file_path)

        if filename not in existing_files:
            missing.append({
                "no_perjanjian": extract_no_perjanjian(filename)
            })

    return missing

def check_sertif_missing_fast():

    SERTIF_DIR = os.path.join(settings.MEDIA_ROOT, "sertifikat")
    SERTIF_DIR_OC = os.path.join(settings.MEDIA_ROOT, "sertifikat_oc")
    SERTIF_DIR_MU = os.path.join(settings.MEDIA_ROOT, "sertifikat_mu")
    SERTIF_DIR_PF = os.path.join(settings.MEDIA_ROOT, "sertifikat_pf")

    result = {
        "LegalDocuments": check_missing_files(LegalDocuments, "certificate_path", SERTIF_DIR),
        "LegalDocumentsOC": check_missing_files(LegalDocumentsOC, "certificate_path", SERTIF_DIR_OC),
        "LegalDocumentsMU": check_missing_files(LegalDocumentMU, "certificate_path", SERTIF_DIR_MU),
        "LegalDocumentsPF": check_missing_files(LegalDocumentsPF, "certificate_path", SERTIF_DIR_PF),
    }

    return result

def check_sertif_view(request):
    data = check_sertif_missing_fast()

    total = sum(len(v) for v in data.values())

    return JsonResponse({
        "total_missing": total,
        "data": data
    })
    
def check_akta_missing_fast():

    PNBP_DIR = os.path.join(settings.MEDIA_ROOT, "pnbp")
    PNBP_DIR_OC = os.path.join(settings.MEDIA_ROOT, "pnbp_oc")
    PNBP_DIR_MU = os.path.join(settings.MEDIA_ROOT, "pnbp_mu")
    PNBP_DIR_PF = os.path.join(settings.MEDIA_ROOT, "pnbp_pf")

    result = {
        "LegalDocuments": check_missing_files(LegalDocuments, "pnbp_path", PNBP_DIR),
        "LegalDocumentsOC": check_missing_files(LegalDocumentsOC, "pnbp_path", PNBP_DIR_OC),
        "LegalDocumentsMU": check_missing_files(LegalDocumentMU, "pnbp_path", PNBP_DIR_MU),
        "LegalDocumentsPF": check_missing_files(LegalDocumentsPF, "pnbp_path", PNBP_DIR_PF),
    }

    return result

def check_akta_view(request):
    data = check_akta_missing_fast()
    
    total = sum(len(v) for v in data.values())
    
    return JsonResponse({
        "total_missing": total,
        "data": data
    })
