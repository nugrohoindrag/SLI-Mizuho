from datetime import date
from decimal import Decimal
from types import SimpleNamespace

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl.utils import column_index_from_string

from . import mizuho_reports as rep
from .mizuho_views import validate_bukti_bayar


def _consumer(**kw):
    base = dict(
        branch_id='0015', branch_name='SERANG', branch=None, agreement_number='0015014718-001',
        debtor_name='HAMID', gender='MALE', marital_status='SINGLE', birth_place='BOGOR',
        birth_date=date(2001, 10, 14), occupation='GURU', address='JL. JAMBU 110', rt='001', rw='003',
        subdistrict='TUGU', district='LALADON', city='BOGOR', province='JAWA BARAT',
        id_number='3175226821000051', spb_date=date(2024, 3, 3), spouse_name='-',
        director_gender=None, director_status=None, director_birth_place=None, director_birth_date=None,
        agreement_date=date(2024, 3, 3), principal_loan='278771537', guarantee_value=Decimal('407040000'),
        otr_value='320850000', brand='MITSUBISHI', type='XPANDER', production_year=2022, color='PUTIH',
        chassis_number='MK3NCTTATNJ003311', engine_number='4A91KBD6677', bpkb_number='T045412356',
        name_on_bpkb='SHINTA', receipt_name='-', tenor=48, installment_start_date=date(2024, 4, 3),
        installment_end_date=date(2028, 3, 3), house_number='B 1111 HKE', postal_code='14260',
        phone_number='0852-15557596', vehicle_condition='USED', service_fee=Decimal('70000'),
        pnbp_fee=Decimal('450000'),
    )
    base.update(kw)
    return SimpleNamespace(**base)


def _doc(consumer=None, **kw):
    base = dict(
        consumer=consumer or _consumer(), document_number='101', voucher_code='RESI-1',
        document_date=date(2026, 10, 13), certificate_number='W10.00123', certificate_date=date(2026, 10, 15),
        notary=SimpleNamespace(full_name='Neltje Tewu', short_title=', S.H., M.Kn.', firm_name='x'),
        service_fee=None, pnbp_fee=Decimal('450000.00'),
    )
    base.update(kw)
    return SimpleNamespace(**base)


class LaporanTests(SimpleTestCase):
    def test_row_matches_format_laporan_columns(self):
        row = rep.laporan_row(1, _doc())
        self.assertEqual(row, [
            1, 'SERANG', '0015014718-001', 'HAMID', '101', 'RESI-1', date(2026, 10, 13),
            'Neltje Tewu, S.H., M.Kn.', 'W10.00123', date(2026, 10, 15), Decimal('70000'), Decimal('450000.00'),
        ])

    def test_flat_fee_used_before_invoice(self):
        doc = _doc(_consumer(service_fee=None))
        self.assertIsNone(rep.laporan_row(1, doc)[10])
        self.assertEqual(rep.laporan_row(1, doc, Decimal('70000'))[10], Decimal('70000'))

    def test_doc_service_fee_overrides_consumer(self):
        row = rep.laporan_row(1, _doc(service_fee=Decimal('65000')))
        self.assertEqual(row[10], Decimal('65000'))

    def test_workbook_keeps_template_header_and_formats_dates(self):
        ws = rep.build_laporan_workbook([_doc(), _doc(document_number='102')]).active
        self.assertEqual(ws['A1'].value, 'No.')
        self.assertEqual(ws['G2'].value, 'Akta/Deed Date')
        self.assertEqual(ws['A4'].value, 1)
        self.assertEqual(ws['E5'].value, '102')
        self.assertEqual(ws['G4'].number_format, 'DD-MM-YYYY')
        self.assertIn('A1:A3', {str(r) for r in ws.merged_cells.ranges})


class FeedbackTests(SimpleTestCase):
    def _cell(self, row, col_letter):
        return row[column_index_from_string(col_letter) - 1]

    def test_headers_are_70_columns_with_feedback_in_bf_bi(self):
        self.assertEqual(len(rep.FEEDBACK_HEADERS), 70)
        self.assertEqual(
            rep.FEEDBACK_HEADERS[column_index_from_string('BF') - 1:column_index_from_string('BI')],
            ['NO AKTA', 'AKTA DATE', 'NO SERTIFIKAT', 'SERTIFIKAT DATE'],
        )

    def test_numeric_deed_number_is_written_as_text(self):
        row = rep.feedback_row(_consumer(), _doc(document_number=401))
        self.assertEqual(self._cell(row, 'BF'), '401')

    def test_row_fills_feedback_and_mirrors_debitur_columns(self):
        row = rep.feedback_row(_consumer(), _doc())
        self.assertEqual(self._cell(row, 'BF'), '101')
        self.assertEqual(self._cell(row, 'BG'), date(2026, 10, 13))
        self.assertEqual(self._cell(row, 'BH'), 'W10.00123')
        self.assertEqual(self._cell(row, 'BI'), date(2026, 10, 15))
        h = rep.FEEDBACK_HEADERS
        self.assertEqual(row[h.index('ALAMAT DEBITUR')], 'JL. JAMBU 110')
        self.assertEqual(row[h.index('IDENTITAS NUMBER FIDUSIA')], '3175226821000051')
        self.assertEqual(row[h.index('NO POL')], 'B 1111 HKE')
        self.assertEqual(row[h.index('PASANGAN PEMBERI FIDUSIA')], '')   # '-' dikosongkan
        self.assertEqual(row[h.index('HUTANG POKOK')], Decimal('278771537'))

    def test_workbook_keeps_leading_zeros_as_text(self):
        ws = rep.build_feedback_workbook([_doc()]).active
        self.assertEqual(ws['A2'].value, '0015')
        self.assertEqual(ws['A2'].number_format, '@')


class BuktiBayarValidationTests(SimpleTestCase):
    def test_accepts_pdf(self):
        self.assertIsNone(validate_bukti_bayar(SimpleUploadedFile('bukti.pdf', b'%PDF-1.4 test')))

    def test_rejects_non_pdf(self):
        self.assertIsNotNone(validate_bukti_bayar(None))
        self.assertIsNotNone(validate_bukti_bayar(SimpleUploadedFile('bukti.jpg', b'%PDF-1.4')))
        self.assertIsNotNone(validate_bukti_bayar(SimpleUploadedFile('bukti.pdf', b'MZ fake exe')))
