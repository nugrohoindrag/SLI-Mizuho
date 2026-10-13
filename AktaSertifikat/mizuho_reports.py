"""
Laporan Mizuho dari Daftar Akta & Sertifikat (requirement §9.3).

a. Laporan Mizuho  : mengikuti Format Laporan.xlsx (header 3 baris, bilingual).
b. File feedback   : format order Mizuho 70 kolom, kolom BF–BI (NO AKTA, AKTA DATE,
                     NO SERTIFIKAT, SERTIFIKAT DATE) terisi.
"""
import os
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

LAPORAN_TEMPLATE = os.path.join(os.path.dirname(__file__), 'report_templates', 'format_laporan_mizuho.xlsx')
LAPORAN_HEADER_ROWS = 3
DATE_FMT = 'DD-MM-YYYY'
NUMBER_FMT = '#,##0'


def _d(v):
    """date/datetime -> date (atau None)."""
    if isinstance(v, datetime):
        return v.date()
    return v if isinstance(v, date) else None


def _num(v):
    if v in (None, '', '-'):
        return None
    try:
        return Decimal(str(v).replace(',', '').strip())
    except InvalidOperation:
        return None


def _s(v):
    return '' if v is None else str(v).strip()


def notary_display(notary):
    """Nama notaris + gelar (mis. 'Neltje Tewu, S.H., M.Kn.')."""
    if notary is None:
        return ''
    name = _s(getattr(notary, 'full_name', ''))
    title = _s(getattr(notary, 'short_title', ''))
    if name:
        return f'{name}{title}'
    return _s(getattr(notary, 'firm_name', ''))


def _service_fee(doc, consumer, default_fee=None):
    # Sama seperti invoice: biaya jasa di legal_documents mengalahkan nilai di consumers.
    # consumers.service_fee baru diisi trigger saat invoice dibuat; sebelum itu pakai tarif flat.
    return (_num(getattr(doc, 'service_fee', None))
            or _num(getattr(consumer, 'service_fee', None))
            or default_fee)


def _pnbp_fee(doc, consumer):
    return _num(getattr(doc, 'pnbp_fee', None)) or _num(getattr(consumer, 'pnbp_fee', None))


# ---------------------------------------------------------------------------
# a. Laporan Mizuho
# ---------------------------------------------------------------------------

LAPORAN_DATE_COLS = (7, 10)      # G Tanggal Akta, J Tanggal Sertifikat
LAPORAN_NUMBER_COLS = (11, 12)   # K Biaya Jasa, L Reimbursement PNBP


def laporan_row(no, doc, default_fee=None):
    consumer = doc.consumer
    branch = getattr(consumer, 'branch', None)
    return [
        no,
        _s(getattr(consumer, 'branch_name', '')) or _s(getattr(branch, 'branch_name', '')),
        _s(consumer.agreement_number),
        _s(consumer.debtor_name),
        _s(doc.document_number),
        _s(doc.voucher_code),
        _d(doc.document_date),
        notary_display(doc.notary),
        _s(doc.certificate_number),
        _d(doc.certificate_date),
        _service_fee(doc, consumer, default_fee),
        _pnbp_fee(doc, consumer),
    ]


def build_laporan_workbook(docs, default_fee=None):
    wb = load_workbook(LAPORAN_TEMPLATE)
    ws = wb.active
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font = copy(ws['A1'].font)

    for i, doc in enumerate(docs, start=1):
        r = LAPORAN_HEADER_ROWS + i
        for c, value in enumerate(laporan_row(i, doc, default_fee), start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = border
            cell.font = font
            if c in LAPORAN_DATE_COLS:
                cell.number_format = DATE_FMT
            elif c in LAPORAN_NUMBER_COLS:
                cell.number_format = NUMBER_FMT
    return wb


# ---------------------------------------------------------------------------
# b. File feedback order (70 kolom, urutan = format baku order Mizuho)
# ---------------------------------------------------------------------------

def _pemberi(c):
    """Kolom alamat/identitas pemberi fidusia; dipakai juga untuk grup FIDUSIA & DEBITUR."""
    return {
        'ALAMAT': c.address, 'RT': c.rt, 'RW': c.rw, 'KELURAHAN': c.subdistrict,
        'KECAMATAN': c.district, 'KABUPATEN': c.city, 'PROVINSI': c.province,
        'IDENTITAS NUMBER': c.id_number,
    }


def feedback_values(consumer, doc):
    c = consumer
    p = _pemberi(c)
    spouse = _s(c.spouse_name)
    values = {
        'BRANCH': c.branch_id,
        'NAMA CABANG': c.branch_name,
        'PEMBERI FIDUSIA': c.debtor_name,
        'JENIS KELAMIN PEMBERI FIDUSIA': c.gender,
        'MARITAL STATUS PEMBERI FIDUSIA': c.marital_status,
        'TEMPAT LAHIR': c.birth_place,
        'TANGGAL LAHIR': c.birth_date,
        'PEKERJAAN': c.occupation,
        'IDENTITAS NUMBER PEMBERI FIDUSIA': c.id_number,
        'TANGGAL KUASA': c.spb_date,
        # Kolom TAF dipakai ulang untuk pasangan (import_cabang FIELD_MAPPING)
        'PASANGAN PEMBERI FIDUSIA': '' if spouse == '-' else spouse,
        'JENIS KELAMIN PASANGAN PEMBERI FIDUSIA': c.director_gender,
        'MARITAL STATUS PASANGAN PEMBERI FIDUSIA': c.director_status,
        'TEMPAT LAHIR PASANGAN': c.director_birth_place,
        'TANGGAL LAHIR PASANGAN': c.director_birth_date,
        'NOMOR KONTRAK': c.agreement_number,
        'TANGGAL KONTRAK': c.agreement_date,
        'NAMA DEBITUR': c.debtor_name,
        'JENIS KELAMIN DEBITUR': c.gender,
        'MARITAL STATUS DEBITUR': c.marital_status,
        'TEMPAT LAHIR DEBITUR': c.birth_place,
        'TANGGAL LAHIR DEBITUR': c.birth_date,
        'HUTANG POKOK': _num(c.principal_loan),
        'NILAI JAMINAN': _num(c.guarantee_value),
        'NILAI BARANG': _num(c.otr_value),
        'MERK': c.brand,
        'TYPE': c.type,
        'TAHUN': c.production_year,
        'WARNA': c.color,
        'NOMOR RANGKA': c.chassis_number,
        'NOMOR MESIN': c.engine_number,
        'NOMOR BPKB/ COVER NOTE': c.bpkb_number,
        'NAMA_BPKB': c.name_on_bpkb,
        'ATAS NAMA KUITANSI': c.receipt_name,
        'NO AKTA': _s(getattr(doc, 'document_number', None)),  # teks: jaga nol di depan
        'AKTA DATE': getattr(doc, 'document_date', None),
        'NO SERTIFIKAT': _s(getattr(doc, 'certificate_number', None)),
        'SERTIFIKAT DATE': getattr(doc, 'certificate_date', None),
        'TENOR': c.tenor,
        'MULAI ANGSURAN': c.installment_start_date,
        'BERAKHIR ANGSURAN': c.installment_end_date,
        'NO BPKB': c.bpkb_number,
        'NO POL': c.house_number,
        'NAMA BPKB': c.name_on_bpkb,
        'KODE POS PEMBERI FIDUSIA': c.postal_code,
        'NO HANDPHONE PEMBERI FIDUSIA': c.phone_number,
        'TIPE': c.vehicle_condition,
    }
    for key, val in p.items():
        values.setdefault(key, val)
        values[f'{key} FIDUSIA'] = val
        values[f'{key} DEBITUR'] = val
    return values


FEEDBACK_HEADERS = [
    'BRANCH', 'NAMA CABANG', 'PEMBERI FIDUSIA', 'JENIS KELAMIN PEMBERI FIDUSIA',
    'MARITAL STATUS PEMBERI FIDUSIA', 'TEMPAT LAHIR', 'TANGGAL LAHIR', 'PEKERJAAN', 'ALAMAT', 'RT', 'RW',
    'KELURAHAN', 'KECAMATAN', 'KABUPATEN', 'PROVINSI', 'IDENTITAS NUMBER PEMBERI FIDUSIA', 'TANGGAL KUASA',
    'PASANGAN PEMBERI FIDUSIA', 'JENIS KELAMIN PASANGAN PEMBERI FIDUSIA', 'MARITAL STATUS PASANGAN PEMBERI FIDUSIA',
    'TEMPAT LAHIR PASANGAN', 'TANGGAL LAHIR PASANGAN', 'ALAMAT FIDUSIA', 'RT FIDUSIA', 'RW FIDUSIA',
    'KELURAHAN FIDUSIA', 'KECAMATAN FIDUSIA', 'KABUPATEN FIDUSIA', 'PROVINSI FIDUSIA', 'IDENTITAS NUMBER FIDUSIA',
    'NOMOR KONTRAK', 'TANGGAL KONTRAK', 'NAMA DEBITUR', 'JENIS KELAMIN DEBITUR', 'MARITAL STATUS DEBITUR',
    'TEMPAT LAHIR DEBITUR', 'TANGGAL LAHIR DEBITUR', 'ALAMAT DEBITUR', 'RT DEBITUR', 'RW DEBITUR',
    'KELURAHAN DEBITUR', 'KECAMATAN DEBITUR', 'KABUPATEN DEBITUR', 'PROVINSI DEBITUR', 'IDENTITAS NUMBER DEBITUR',
    'HUTANG POKOK', 'NILAI JAMINAN', 'NILAI BARANG', 'MERK', 'TYPE', 'TAHUN', 'WARNA', 'NOMOR RANGKA',
    'NOMOR MESIN', 'NOMOR BPKB/ COVER NOTE', 'NAMA_BPKB', 'ATAS NAMA KUITANSI', 'NO AKTA', 'AKTA DATE',
    'NO SERTIFIKAT', 'SERTIFIKAT DATE', 'TENOR', 'MULAI ANGSURAN', 'BERAKHIR ANGSURAN', 'NO BPKB', 'NO POL',
    'NAMA BPKB', 'KODE POS PEMBERI FIDUSIA', 'NO HANDPHONE PEMBERI FIDUSIA', 'TIPE',
]

FEEDBACK_MONEY = {'HUTANG POKOK', 'NILAI JAMINAN', 'NILAI BARANG'}


def feedback_row(consumer, doc):
    values = feedback_values(consumer, doc)
    row = []
    for h in FEEDBACK_HEADERS:
        v = values.get(h)
        if isinstance(v, str):
            v = v.strip()
            v = '' if v == '-' else v
        row.append(_d(v) if isinstance(v, (date, datetime)) else v)
    return row


def build_feedback_workbook(docs):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(FEEDBACK_HEADERS)
    for doc in docs:
        ws.append(feedback_row(doc.consumer, doc))
    for col_idx, h in enumerate(FEEDBACK_HEADERS, start=1):
        for (cell,) in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            if isinstance(cell.value, date):
                cell.number_format = DATE_FMT
            elif h in FEEDBACK_MONEY:
                cell.number_format = NUMBER_FMT
            elif isinstance(cell.value, str):
                cell.number_format = '@'  # jaga nol di depan (no. kontrak, NIK, kode pos)
    return wb
