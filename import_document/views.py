from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from rest_framework.response import Response
from rest_framework import viewsets, permissions, status as drf_status
from rest_framework.authentication import SessionAuthentication
from rest_framework.decorators import api_view, permission_classes, parser_classes, action
from openpyxl import Workbook
from rest_framework.parsers import MultiPartParser, FormParser
from django.utils import timezone
import os, posixpath, logging, openpyxl
from django.conf import settings
from django.core.files.storage import default_storage
from django.utils.crypto import get_random_string
from django.db.models import Q, OuterRef, Subquery, Count, Exists, Case, When, F, CharField, Value, Window
from django.db.models.functions import RowNumber
from django.http import HttpResponse, HttpResponseBadRequest, Http404, FileResponse, HttpResponseForbidden
from rest_framework.permissions import IsAuthenticated
from django.core.exceptions import FieldDoesNotExist
from urllib3 import request
from import_cabang.models import Consumer, Branch
from .models import ConsumerDocument
from .serializers import ConsumerDocumentSerializer
from AdminMenu.auth_helpers import (
    permission_required,
    class_permission_required,
    get_user_permissions,
    get_user_role_names,
    user_has_permission,
)
from common.const.role_map import has_role, ROLE_KEYS
from common.validators import validate_excel_file_or_error
from AdminUser.branch_access import (
    branch_scope_q,
    filter_queryset_by_user_branches,
    get_primary_user_branch_code,
    get_user_branch_codes,
    user_is_branch_restricted,
)
from rest_framework.pagination import CursorPagination
from datetime import datetime, date, timedelta, time
from django.db import transaction, models, connection
from django.utils.text import slugify
from django.core.paginator import Paginator
from rest_framework.exceptions import ValidationError
from io import BytesIO
from django.urls import reverse
from django.views.decorators.clickjacking import xframe_options_sameorigin
from psycopg2 import sql

logger = logging.getLogger(__name__)

# ===== HTML page view (non-API) =====

STATUS_MAP = {
    'uncomplete': 0,
    'complete':   1,
    'process':    2,
    'revisi':     3,
    'need revisi': 3,
}

@login_required(login_url='')
@permission_required('read-dokumen')
def view_staff_cabang_document(request):
    branches = Branch.objects.all().order_by('branch_name')

    if _user_is_staff_cabang(request.user):
        key = _consumer_branch_filter_kwargs()
        user_branches = get_user_branch_codes(request.user)
        if not user_branches or not key:
            base = Consumer.objects.none()
        else:
            base = Consumer.objects.filter(branch_scope_q(key, user_branches))
    else:
        base = Consumer.objects.all()

    # hitung consumer_name berdasar fiducia_provider
    qs = (
        base.annotate(
            consumer_name=Case(
                # Kasus "Atas Nama BPKB" atau "BPKB"
                When(fiducia_provider__iexact='Atas Nama BPKB', then=F('name_on_bpkb')),
                When(fiducia_provider__iexact='BPKB', then=F('name_on_bpkb')),

                # Kasus "Atas Nama Customer" atau "CUSTOMER"
                When(fiducia_provider__iexact='Atas Nama Customer', then=F('debtor_name')),
                When(fiducia_provider__iexact='CUSTOMER', then=F('debtor_name')),

                # Default fallback → debtor_name
                default=F('debtor_name'),
                output_field=CharField(),
            )
        )
        .only(
            'id',
            'agreement_number',
            'debtor_name',
            'upload_date',
            'branch_name',
            'fiducia_provider',
            'name_on_bpkb',
        )
        .order_by('-id')
    )

    try:
        page_size = min(int(request.GET.get('page_size', 25)), 200)
    except (TypeError, ValueError):
        page_size = 25

    last_id = request.GET.get('last_id')
    if last_id:
        qs = qs.filter(id__lt=last_id)

    items = list(qs[:page_size])
    next_last_id = items[-1].id if items else None

    user_permissions = list(get_user_permissions(request.user))
    user_roles = ['Warkah', 'superadministrator'] if request.user.is_superuser else list(get_user_role_names(request.user))

    return render(
        request,
        'staff_cabang/page-staff_cabang-dokument.html',
        {
            'branches': branches,
            'consumer_list': items,
            'next_last_id': next_last_id,
            'page_size': page_size,
            'user_permissions': user_permissions,
            'user_roles': user_roles,
        }
    )

def can_view_all(user):
    role_names = []
    if getattr(user, 'roles', None) is not None:
        role_names = list(user.roles.values_list('name', flat=True))
    return (
        has_role(role_names, ROLE_KEYS["SUPERADMIN"])
        or has_role(role_names, ROLE_KEYS["HO"])
        or has_role(role_names, ROLE_KEYS["OPERATOR"])
        or has_role(role_names, ROLE_KEYS["WARKAH"])
    )

def _user_is_staff_cabang(user):
    """Check if user has staff_cabang role"""
    return user_is_branch_restricted(user)

def get_user_branch_code(user):
    return get_primary_user_branch_code(user)

def _consumer_branch_filter_kwargs():
    """
    Kembalikan kwargs filter untuk Consumer berdasarkan struktur model:
    - Jika ada FK 'branch' → gunakan branch__branch_code
    - Jika ada kolom 'branch_code' → gunakan branch_code
    """
    try:
        Consumer._meta.get_field('branch')      # FK
        return 'branch__branch_code'
    except FieldDoesNotExist:
        try:
            Consumer._meta.get_field('branch_code')
            return 'branch_code'
        except FieldDoesNotExist:
            return None

ALLOWED_EXT  = {'.pdf', '.jpg', '.jpeg', '.png'}
ALLOWED_MIME = {'application/pdf', 'image/jpeg', 'image/png'}

def _get_ext(filename: str) -> str:
    _, ext = os.path.splitext(filename or '')
    return (ext or '.bin').lower()

def _is_allowed_upload(upload) -> bool:
    ext_ok = _get_ext(getattr(upload, 'name', '')) in ALLOWED_EXT
    mime = getattr(upload, 'content_type', '') or ''
    return ext_ok and mime in ALLOWED_MIME

def _max_bytes_for_consumer(consumer) -> int:
    """
    Batas ukuran file:
      - Corporate: 10 MB  (kode 'C' atau 'PT')
      - Individual: 5 MB  (kode 'I' atau default lain)
    """
    ctype = ((getattr(consumer, 'customer_type', '') or '').upper()).strip()

    TEN_MB = 10 * 1024 * 1024
    FIVE_MB = 5 * 1024 * 1024

    if ctype in {'C', 'PT', 'CORP', 'CORPORATE', 'COMPANY'}:
        return TEN_MB
    if ctype in {'I', 'INDIVIDU', 'INDIVIDUAL'}:
        return FIVE_MB

    return FIVE_MB

@xframe_options_sameorigin
@login_required
def protected_media_view(request, path):
    agreement_no = _extract_agreement_no(path)
    _authorize_media_request(
        request,
        agreement_no,
        table_name="consumers",
        permission_name="read-dokumen",
    )
    return _serve_media_file(path)

@xframe_options_sameorigin
@login_required
def protected_media_view_mu(request, path):
    agreement_no = _extract_agreement_no(path)
    _authorize_media_request(
        request,
        agreement_no,
        table_name="consumers_mu",
        permission_name="read-dokumen-mu",
    )
    return _serve_media_file(path)

@xframe_options_sameorigin
@login_required
def protected_media_view_oc(request, path):
    agreement_no = _extract_agreement_no(path)
    _authorize_media_request(
        request,
        agreement_no,
        table_name="consumers_oc",
        permission_name="read-dokumen-oc",
    )
    return _serve_media_file(path)

@xframe_options_sameorigin
@login_required
def protected_media_view_pf(request, path):
    agreement_no = _extract_agreement_no(path)
    _authorize_media_request(
        request,
        agreement_no,
        table_name="change_of_fiduciary",
        permission_name="read-dokumen-pf",
    )
    return _serve_media_file(path)


def _extract_agreement_no(path: str):
    parts = [p for p in (path or "").split("/") if p]
    if len(parts) < 2:
        return None
    return parts[1]


def _get_user_branch_code(user):
    return get_primary_user_branch_code(user)


def _lookup_branch_code(table_name: str, agreement_no: str):
    # Whitelist valid table names to prevent SQL injection
    ALLOWED_TABLES = {
        "consumers",
        "consumers_mu",
        "consumers_oc",
        "change_of_fiduciary",
    }

    if table_name not in ALLOWED_TABLES:
        raise ValueError("Invalid table name")
    query = sql.SQL(
        "SELECT branch_code FROM {} WHERE agreement_number = %s LIMIT 1"
    ).format(sql.Identifier(table_name))

    with connection.cursor() as cursor:
        cursor.execute(query, [agreement_no])
        row = cursor.fetchone()
    return row[0] if row else None


def _authorize_media_request(request, agreement_no, table_name, permission_name):
    user = request.user

    if not agreement_no:
        raise Http404("File not found")

    if not (user.is_superuser or user_has_permission(user, permission_name)):
        raise Http404("File not found")

    if _user_is_staff_cabang(user):
        user_branches = get_user_branch_codes(user)
        if not user_branches:
            raise Http404("File not found")

        branch_code = _lookup_branch_code(table_name, agreement_no)
        if not branch_code:
            raise Http404("File not found")

        if branch_code not in user_branches:
            raise Http404("File not found")


def _serve_media_file(path: str):
    media_root = os.path.abspath(settings.MEDIA_ROOT)
    file_path = os.path.abspath(os.path.join(media_root, path))

    if not file_path.startswith(media_root + os.sep) or not os.path.exists(file_path):
        raise Http404("File not found")

    return FileResponse(open(file_path, "rb"))

def _build_public_url(path: str, request=None) -> str:
    if not path:
        return ''
    return reverse('media', kwargs={'path': path})

def _storage_path(agreement_number: str, filename: str) -> str:
    ts   = timezone.now().strftime('%Y%m%d%H%M%S')
    base = slugify((filename or '').rsplit('.', 1)[0]) or 'upload'
    ext  = _get_ext(filename)
    folder = slugify(agreement_number or 'unknown') or 'unknown'
    return f"consumer_documents/{folder}/{ts}_{base}{ext}"

class ConsumerDocumentViewSet(viewsets.ModelViewSet):
    authentication_classes = [SessionAuthentication]
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    serializer_class = ConsumerDocumentSerializer
    queryset = ConsumerDocument.objects.all()
    
    class DocCursorPagination(CursorPagination):
        page_size = 50
        page_size_query_param = 'page_size'
        max_page_size = 200
        ordering = '-id'

    pagination_class = DocCursorPagination

    # ===================== SCOPE & FILTER =====================
    def _agreement_numbers_for_user_branch(self):
        """Ambil agreement_number yang boleh dilihat user, berdasar branch_code user."""
        if not _user_is_staff_cabang(self.request.user):
            return Consumer.objects.all().values_list('agreement_number', flat=True)

        user_branches = get_user_branch_codes(self.request.user)
        if not user_branches:
            return Consumer.objects.none().values_list('agreement_number', flat=True)

        return Consumer.objects.filter(branch__branch_code__in=user_branches).values_list('agreement_number', flat=True)

    def get_queryset(self):
        user = self.request.user
        qs = super().get_queryset().only('id', 'agreement_number', 'consumer_name', 'updated_at', 'file_url')
        user_branches = get_user_branch_codes(user)
        if _user_is_staff_cabang(user) and user_branches:
            exists_q = Consumer.objects.filter(
                agreement_number=OuterRef('agreement_number'),
                branch__branch_code__in=user_branches
            )
            qs = qs.filter(Exists(exists_q))
        elif _user_is_staff_cabang(user):
            qs = qs.none()
        return qs

    def _build_consumer_map(self, agreement_numbers):
        consumer_qs = (
            Consumer.objects.filter(agreement_number__in=agreement_numbers).select_related('branch').only('agreement_number', 'debtor_name', 'customer_type', 'status', 'document_status','name_on_bpkb', 'fiducia_provider', 'agreement_date', 'branch__branch_code', 'branch__branch_name'))
        return {c.agreement_number: c for c in consumer_qs}

    def get_serializer_context(self):
        context = super().get_serializer_context()
        # Default: KOSONG (jangan eager load semuanya)
        context['consumer_map'] = {}
        # Kalau ada daftar terbatas untuk halaman saat ini, baru build map
        agreement_numbers = getattr(self.request, "_page_agreements", None)
        if agreement_numbers:
            context['consumer_map'] = self._build_consumer_map(agreement_numbers)
        return context

    # ===================== LIST & LOOKUP =====================
    @class_permission_required("read-dokumen")
    def list(self, request, *args, **kwargs):
        draw = request.query_params.get('draw')
        if draw is None:
            return super().list(request, *args, **kwargs)

        start  = int(request.query_params.get('start', 0))
        length = int(request.query_params.get('length', 10))
        search_value = (request.query_params.get('search[value]') or '').strip()

        base_qs = self.get_queryset()

        from datetime import datetime, date
        DATE_FIELD = 'created_at'

        def _parse_date(s: str):
            if not s: return None
            s = s.strip().replace(' ', '')
            for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
                try:
                    return datetime.strptime(s, fmt).date()
                except ValueError:
                    pass
            return None

        from_d = _parse_date(request.query_params.get('fromDate') or '')
        to_d   = _parse_date(request.query_params.get('toDate') or '')
        include_old = (request.query_params.get('include_old') or '').lower() in ('1', 'true', 'yes')

        if not from_d and not to_d:
            if search_value or include_old:
                from_d = None
                to_d = None
            else:
                today = date.today()
                from_d = to_d = today
        elif from_d and not to_d:
            to_d = from_d
        elif to_d and not from_d:
            from_d = to_d
        
        if from_d and to_d:
            diff_days = (to_d - from_d).days + 1

            if diff_days > 31:
                return Response(
                    {
                        "error": "Filter tanggal maksimal 31 hari (inklusif).",
                    },
                    status=400
                )

        qs = base_qs

        if from_d or to_d:
            consumer_date_q = Consumer.objects.filter(agreement_number=OuterRef('agreement_number'))
            if from_d:
                consumer_date_q = consumer_date_q.filter(**{f'{DATE_FIELD}__date__gte': from_d})
            if to_d:
                consumer_date_q = consumer_date_q.filter(**{f'{DATE_FIELD}__date__lte': to_d})
            qs = qs.filter(Exists(consumer_date_q))

        if search_value:
            qs = qs.filter(
                Q(agreement_number__icontains=search_value) |
                Q(consumer_name__icontains=search_value)
            )

        records_total = (base_qs.values('agreement_number').annotate(n=Count('id')).count())
        records_filtered = (qs.values('agreement_number').annotate(n=Count('id')).count())

        order_col_idx = request.query_params.get('order[0][column]')
        order_dir = request.query_params.get('order[0][dir]', 'asc')
        order_field_from_dt = request.query_params.get(f'columns[{int(order_col_idx)}][data]') if order_col_idx else None

        ALLOWED_ORDER_FIELDS = {
            'agreement_number': 'agreement_number',
            'updated_at': 'updated_at',
            'file_url': 'file_url',
        }
        order_field = ALLOWED_ORDER_FIELDS.get(order_field_from_dt, 'updated_at')
        order_by = F(order_field).desc(nulls_last=True) if order_dir == 'desc' else F(order_field).asc(nulls_last=True)

        partition_fields = [F('agreement_number')]
        qs_one = qs.annotate(
            rn=Window(
                expression=RowNumber(),
                partition_by=partition_fields,
                order_by=[order_by]
            )
        ).filter(rn=1)

        ordered = qs_one.order_by(order_by, F('agreement_number').asc(nulls_last=True), F('id').desc())
        page_qs = ordered[start:start+length]
        self.request._page_agreements = list(page_qs.values_list('agreement_number', flat=True))

        serializer = self.get_serializer(page_qs, many=True)
        return Response({
            'draw': int(draw),
            'recordsTotal': records_total,
            'recordsFiltered': records_filtered,
            'data': serializer.data,
        })

    @action(detail=False, methods=['get'], url_path='customer-type')
    @class_permission_required("read-dokumen")
    def customer_type_lookup(self, request):
        an = (request.query_params.get('agreement_number') or '').strip()
        if not an:
            return Response({'detail': 'agreement_number wajib diisi.'}, status=drf_status.HTTP_400_BAD_REQUEST)

        q = Consumer.objects.filter(agreement_number=an).select_related('branch')
        if _user_is_staff_cabang(request.user):
            user_branches = get_user_branch_codes(request.user)
            if user_branches:
                q = q.filter(branch__branch_code__in=user_branches)
            else:
                q = q.none()

        consumer = q.first()
        if not consumer:
            # 404 kalau nomor tidak ada atau tidak berhak
            return Response({'detail': 'Nomor perjanjian tidak ditemukan atau tidak berhak diakses.'},
                            status=drf_status.HTTP_404_NOT_FOUND)

        return Response({
            'agreement_number': an,
            'customer_type': (getattr(consumer, 'customer_type', None) or '').upper()
        })

    # ===================== CREATE / BULK =====================
    @action(detail=False, methods=['post'], url_path='batch-create')
    @class_permission_required("create-dokumen")
    def batch_create(self, request):
        data_list = request.data if isinstance(request.data, list) else [request.data]
        serializer = self.get_serializer(data=data_list, many=True)
        serializer.is_valid(raise_exception=True)
        response_serializer = self.perform_bulk_create(serializer)
        return Response(response_serializer.data, status=drf_status.HTTP_201_CREATED)

    def perform_bulk_create(self, serializer):
        serializer.save()
        agreement_numbers = [obj.agreement_number for obj in serializer.instance]
        consumer_map = self._build_consumer_map(agreement_numbers)
        return self.get_serializer(serializer.instance, many=True, context={'consumer_map': consumer_map})

    @class_permission_required("create-dokumen")
    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        instance = serializer.save()

        consumer = Consumer.objects.filter(agreement_number=instance.agreement_number).select_related('branch').first()
        consumer_map = {instance.agreement_number: consumer} if consumer else {}
        return Response(self.get_serializer(instance, context={'consumer_map': consumer_map}).data,status=drf_status.HTTP_201_CREATED)

    # ===================== PATCH (UPLOAD FILE) =====================
    @class_permission_required("update-dokumen")
    def partial_update(self, request, *args, **kwargs):
        instance: ConsumerDocument = self.get_object()

        # ==== 1) Cek scope cabang (tetap) ====
        if _user_is_staff_cabang(request.user):
            user_branches = get_user_branch_codes(request.user)
            consumer_for_scope = Consumer.objects.filter(
                agreement_number=instance.agreement_number
            ).select_related('branch').only('branch__branch_code').first()
            doc_branch = getattr(getattr(consumer_for_scope, 'branch', None), 'branch_code', None)
            if not user_branches or not doc_branch or doc_branch not in user_branches:
                return Response({'detail': 'Anda tidak berhak mengubah data cabang lain.'},
                                status=drf_status.HTTP_403_FORBIDDEN)

        # ==== 2) Pisahkan payload non-file ====
        blocked_keys = {'branch_code', 'branch_name', 'file'}
        data = {k: request.data.get(k) for k in request.data.keys() if k not in blocked_keys}

        uploaded = request.FILES.get('file')
        saved_path = None
        public_url = None

        # ==== 3) VALIDASI & SIMPAN FILE DI LUAR TRANSAKSI ====
        if uploaded:
            if not _is_allowed_upload(uploaded):
                return Response({'detail': 'Format file harus PDF/JPG/PNG.'}, status=drf_status.HTTP_400_BAD_REQUEST)

            consumer = Consumer.objects.filter(agreement_number=instance.agreement_number)\
                                    .only('status', 'document_status').first()
            if not consumer:
                return Response({'detail': 'Consumer untuk agreement ini tidak ditemukan.'},
                                status=drf_status.HTTP_404_NOT_FOUND)

            max_bytes = _max_bytes_for_consumer(consumer)
            if uploaded.size > max_bytes:
                mb = max_bytes // (1024 * 1024)
                return Response({'detail': f'Ukuran file maksimal {mb} MB.'},
                                status=drf_status.HTTP_400_BAD_REQUEST)

            # I/O file DI LUAR atomic:
            path = _storage_path(instance.agreement_number, uploaded.name)
            saved_path = default_storage.save(path, uploaded)  # bisa lambat
            public_url = _build_public_url(saved_path, request)

        # ==== 4) Update DB dalam transaksi kecil ====
        with transaction.atomic():
            update_fields = []

            if uploaded:
                instance.file_url = public_url
                if hasattr(instance, 'file_path'):
                    instance.file_path = saved_path
                    update_fields.append('file_path')
                update_fields.append('file_url')

                # [OPSI A - Aman] Set created_at hanya jika belum ada (NULL/None)
                if hasattr(instance, 'created_at') and not getattr(instance, 'created_at', None):
                    from django.utils import timezone
                    instance.created_at = timezone.now()
                    update_fields.append('created_at')

                # [OPSI B - Force] Kalau mau created_at SELALU di-set ke waktu upload (overwrite):
                # from django.utils import timezone
                # instance.created_at = timezone.now()
                # update_fields.append('created_at')

                # Sinkronisasi status consumer (singkat & aman)
                consumer = Consumer.objects.filter(agreement_number=instance.agreement_number).first()
                if consumer:
                    set_fields = {'status': 'process'}
                    if hasattr(consumer, 'document_status'):
                        set_fields['document_status'] = STATUS_MAP['process']
                        consumer.status = 'process'
                        consumer.document_status = STATUS_MAP['process']
                        consumer.save(update_fields=list(set_fields.keys()))
                    else:
                        consumer.status = 'process'
                        consumer.save(update_fields=['status'])

                if hasattr(instance, 'modified_by'):
                    instance.modified_by = request.user
                    update_fields.append('modified_by')
                if hasattr(instance, 'document_status'):
                    instance.document_status = STATUS_MAP['process']
                    update_fields.append('document_status')

            if update_fields:
                instance.save(update_fields=update_fields)

            # Patch kolom lain (tanpa memicu context berat)
            serializer = self.get_serializer(instance, data=data, partial=True, context={'consumer_map': {}})
            serializer.is_valid(raise_exception=True)
            instance = serializer.save()

        # ==== 5) Balikan payload dengan context KECIL khusus objek ====
        consumer_ctx = Consumer.objects.filter(agreement_number=instance.agreement_number)\
                                    .select_related('branch')\
                                    .only('agreement_number', 'debtor_name', 'customer_type',
                                            'status', 'document_status', 'name_on_bpkb', 'fiducia_provider',
                                            'agreement_date', 'branch__branch_code', 'branch__branch_name')\
                                    .first()

        payload = self.get_serializer(
            instance,
            context={'consumer_map': {instance.agreement_number: consumer_ctx} if consumer_ctx else {}}
        ).data

        if 'file_url' not in payload:
            payload['file_url'] = getattr(instance, 'file_url', '')

        return Response(payload)

        # except Exception as e:
        #     return Response({'detail': f'{e.__class__.__name__}: {e}'},status=drf_status.HTTP_400_BAD_REQUEST)

@login_required
@permission_required("read-dokumen")
def download_template(request):
    status_raw = request.GET.get('status') or 'process'
    raw_from = (
        request.GET.get('from')
        or request.GET.get('fromDate')
        or request.GET.get('start')
        or request.GET.get('startDate')
    )
    raw_to = (
        request.GET.get('to')
        or request.GET.get('toDate')
        or request.GET.get('end')
        or request.GET.get('endDate')
    )

    ACCEPT_FMTS = [
        "%Y-%m-%d", "%Y/%m/%d",
        "%d/%m/%Y", "%d-%m-%Y",
        "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S",
    ]

    def parse_to_date(s):
        if not s:
            return None
        s = s.strip()
        for fmt in ACCEPT_FMTS:
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        # fallback ambil 10 char pertama
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(s[:10], fmt).date()
            except Exception:
                pass
        return None

    start_date = parse_to_date(raw_from)
    end_date   = parse_to_date(raw_to)

    if not (start_date and end_date):
        return HttpResponse(
            "Bad request: from/to date required.",
            status=400,
            content_type="text/plain",
        )

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    range_days_inclusive = (end_date - start_date).days + 1
    if range_days_inclusive > 7:
        return HttpResponse(
            "Bad request: date range must be ≤ 7 days.",
            status=400,
            content_type="text/plain",
        )

    def _norm(s):
        return '' if s is None else str(s).strip().lower()

    norm_status = _norm(status_raw) or 'process'

    start_dt = datetime.combine(start_date, time.min)
    end_dt   = datetime.combine(end_date, time.max)

    consumers_qs = Consumer.objects.filter(
        status__iexact=norm_status,
        created_at__range=(start_dt, end_dt),
    )
    if _user_is_staff_cabang(request.user):
        key = _consumer_branch_filter_kwargs()
        user_branches = get_user_branch_codes(request.user)
        if not user_branches or not key:
            return HttpResponseForbidden("Branch tidak ditemukan untuk user ini.")
        consumers_qs = consumers_qs.filter(branch_scope_q(key, user_branches))

    consumers = list(
        consumers_qs.values('agreement_number', 'status').order_by('agreement_number')
    )

    if not consumers:
        return HttpResponse(
            "No data for given status/date range.",
            status=204,
            content_type="text/plain",
        )

    ag_nums = [c['agreement_number'] for c in consumers if c['agreement_number']]

    docs_qs = (
        ConsumerDocument.objects
        .filter(agreement_number__in=ag_nums)
        .order_by('agreement_number', '-id')
        .values('agreement_number', 'notes', 'modified_by__username')
    )

    latest_doc_map = {}
    for d in docs_qs:
        ag = d['agreement_number']
        if ag not in latest_doc_map:
            latest_doc_map[ag] = {
                'notes': d['notes'] or '',
                'uploader': d['modified_by__username'] or '',
            }

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Template")

    ws.append(["Notes: Isi Status Dengan Revisi/Complete"])
    ws.append([])

    headers = ['No Perjanjian', 'Status', 'Upload_by', 'Notes']
    ws.append(headers)

    for c in consumers:
        ag = c['agreement_number']
        status = c['status'] or ''
        doc_info = latest_doc_map.get(ag, {})
        uploader = doc_info.get('uploader', '')
        notes    = doc_info.get('notes', '')

        ws.append([
            ag or '',
            status,
            uploader,
            notes,
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    status_suffix = (norm_status or 'process').replace(' ', '_')
    if start_date == end_date:
        suffix = f"{status_suffix}_{start_date.strftime('%Y%m%d')}"
    else:
        suffix = f"{status_suffix}_{start_date.strftime('%Y%m%d')}-{end_date.strftime('%Y%m%d')}"

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=template_{suffix}.xlsx'
    return response

def _normalize_status(val):
    """
    Normalisasi input status ke lower-case tanpa spasi berlebih.
    Silakan gunakan versi Anda bila sudah ada.
    """
    if val is None:
        return ''
    return str(val).strip().lower()

@api_view(['POST'])
@parser_classes([MultiPartParser, FormParser])
@permission_classes([IsAuthenticated])
@permission_required("update-dokumen")
def upload_template_excel(request):
    excel_file = request.FILES.get('excel_file')

    if not excel_file:
        return Response({'detail': 'No file uploaded.'}, status=drf_status.HTTP_400_BAD_REQUEST)

    # Validate Excel file
    validation_error = validate_excel_file_or_error(excel_file)
    if validation_error:
        return validation_error

    try:
        import openpyxl  # pastikan sudah terpasang
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))

        # cari header secara dinamis
        header_row_idx = None
        for idx, row in enumerate(rows):
            if row and row[0] and str(row[0]).strip() == 'No Perjanjian':
                header_row_idx = idx
                break

        if header_row_idx is None:
            return Response({'detail': 'Header not found in file.'}, status=drf_status.HTTP_400_BAD_REQUEST)

        header = rows[header_row_idx]
        data_rows = rows[header_row_idx + 1:]

        expected_header = ['No Perjanjian', 'Status', 'Upload_by', 'Notes']
        if [(h.strip() if h else '') for h in header[:4]] != expected_header:
            return Response({'detail': 'Invalid template format.'}, status=drf_status.HTTP_400_BAD_REQUEST)

        updated_count = 0
        not_found_list = []

        # Kerjakan dalam transaksi supaya konsisten
        with transaction.atomic():
            for row in data_rows:
                if not row or all(v is None for v in row):
                    continue  # skip baris kosong

                agreement_number, status_val, _, notes = row[:4]

                if not agreement_number or not status_val:
                    continue

                # normalisasi agreement_number ke string
                if isinstance(agreement_number, float):
                    agreement_number = str(int(agreement_number)).strip()
                elif isinstance(agreement_number, int):
                    agreement_number = str(agreement_number).strip()
                else:
                    agreement_number = str(agreement_number).strip()

                # normalisasi status -> map ke kode integer
                norm_status = _normalize_status(status_val)
                mapped_doc_status = STATUS_MAP.get(norm_status, STATUS_MAP['uncomplete'])

                consumer = Consumer.objects.filter(agreement_number=agreement_number).first()

                if consumer:
                    # Jika punya field 'status' tekstual, tetap update (opsional)
                    if hasattr(consumer, 'status'):
                        consumer.status = norm_status

                    # === INTI PERMINTAAN: set document_status pakai STATUS_MAP ===
                    consumer.document_status = mapped_doc_status
                    # simpan field yang ada saja
                    fields_to_update = ['document_status']
                    if hasattr(consumer, 'status'):
                        fields_to_update.append('status')
                    consumer.save(update_fields=fields_to_update)

                    # Update / create ConsumerDocument
                    consumer_doc = ConsumerDocument.objects.filter(agreement_number=agreement_number).first()
                    if consumer_doc:
                        # kalau model punya field 'consumer_status' (text), ikutkan update
                        if hasattr(consumer_doc, 'consumer_status'):
                            consumer_doc.consumer_status = norm_status
                        consumer_doc.modified_by = request.user
                        consumer_doc.notes = notes
                        consumer_doc.updated_at = timezone.now()
                        consumer_doc.save()
                    else:
                        kwargs = dict(
                            agreement_number=agreement_number,
                            consumer_name=getattr(consumer, 'debtor_name', '') or getattr(consumer, 'consumer_name', ''),
                            agreement_date=getattr(consumer, 'agreement_date', None),
                            document_type='Default Type',
                            description='Auto-generated document',
                            notes=notes,
                            file_url='',
                            modified_by=request.user,
                            created_at=timezone.now(),
                            updated_at=timezone.now()
                        )
                        # jika ada field 'consumer_status' di ConsumerDocument (text)
                        if 'consumer_status' in [f.name for f in ConsumerDocument._meta.get_fields()]:
                            kwargs['consumer_status'] = norm_status
                        ConsumerDocument.objects.create(**kwargs)

                    updated_count += 1
                else:
                    not_found_list.append(agreement_number)

        msg = f"{updated_count} consumer updated."
        if not_found_list:
            not_found_str = ', '.join(str(x) for x in not_found_list)
            msg += f" Agreement not found: {not_found_str}"

        return Response({'detail': msg}, status=drf_status.HTTP_200_OK)

    except Exception as e:
        return Response({'detail': f'Error processing file: {str(e)}'}, status=drf_status.HTTP_400_BAD_REQUEST)

@api_view(['POST'])
@parser_classes([MultiPartParser, FormParser])
@permission_classes([IsAuthenticated])
@permission_required("update-dokumen")
def upload_template_patch(request):
    excel_patch = request.FILES.get('excel_patch')

    if not excel_patch:
        return Response({'detail': 'No file uploaded.'}, status=drf_status.HTTP_400_BAD_REQUEST)

    validation_error = validate_excel_file_or_error(excel_patch)
    if validation_error:
        return validation_error

    try:
        wb = openpyxl.load_workbook(excel_patch)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))

        header_row_idx = None
        for idx, row in enumerate(rows):
            if row and row[0] and str(row[0]).strip() == 'No Perjanjian' and 'Patch File' in row:
                header_row_idx = idx
                break

        if header_row_idx is None:
            return Response({'detail': 'Cek Header Anda [No Perjanjian, Patch File]'}, status=drf_status.HTTP_400_BAD_REQUEST)

        header = rows[header_row_idx]
        data_rows = rows[header_row_idx + 1:]

        expected_header = ['No Perjanjian', 'Patch File']
        if [(h.strip() if h else '') for h in header[:2]] != expected_header:
            return Response({'detail': 'Invalid template format.'}, status=drf_status.HTTP_400_BAD_REQUEST)

        updated_count = 0
        not_found_list = []

        with transaction.atomic():
            for row in data_rows:
                if not row or not row[0]:
                    continue

                no_perjanjian = str(row[0]).strip()
                file_name = str(row[1]).strip() if row[1] else None

                if not file_name:
                    not_found_list.append(f"{no_perjanjian} (filename kosong)")
                    continue

                try:
                    doc = ConsumerDocument.objects.filter(
                        agreement_number=no_perjanjian
                    ).first()

                    if not doc:
                        not_found_list.append(no_perjanjian)
                        continue

                    new_path = f"/media/consumer_documents/{file_name}"

                    doc.file_url = new_path
                    doc.save()

                    updated_count += 1

                except Exception as e:
                    not_found_list.append(no_perjanjian)
                

        msg = f"{updated_count} Konsumen Path updated."
        if not_found_list:
            not_found_str = ', '.join(str(x) for x in not_found_list)
            msg += f" Agreement not found: {not_found_str}"

        return Response({'detail': msg}, status=drf_status.HTTP_200_OK)

    except Exception as e:
        return Response({'detail': f'Error processing file: {str(e)}'}, status=drf_status.HTTP_400_BAD_REQUEST)
